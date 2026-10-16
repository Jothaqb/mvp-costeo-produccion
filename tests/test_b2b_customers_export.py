import asyncio
import hashlib
import io
import os
import tempfile
import unittest
from datetime import date, datetime, timedelta, timezone
from decimal import Decimal
from pathlib import Path
from unittest import mock

from fastapi import HTTPException
from openpyxl import load_workbook
from sqlalchemy import event


_PREVIOUS_DATABASE_URL = os.environ.get("DATABASE_URL")
_TEMP_DB = tempfile.NamedTemporaryFile(prefix="b2b_customers_export_", suffix=".db", delete=False)
_TEMP_DB.close()
os.environ["DATABASE_URL"] = f"sqlite:///{Path(_TEMP_DB.name).as_posix()}"

with mock.patch("zoneinfo.ZoneInfo", return_value=timezone.utc):
    from app import main  # noqa: E402
from app.models import (  # noqa: E402
    B2BAROpeningBalance,
    B2BARPayment,
    B2BCommercialAgreement,
    B2BCustomer,
    B2BSalesOrder,
    B2BSalesOrderLine,
    B2CSalesOrder,
    Channel,
    Role,
    User,
    UserRole,
    UserSession,
)
from app.services.accounts_receivable_service import (  # noqa: E402
    build_accounts_receivable_dashboard,
)
from app.services.b2b_customer_export_service import (  # noqa: E402
    B2B_CUSTOMER_EXPORT_HEADERS,
    NO_REGISTERED_SIGNATURE,
    b2b_customer_export_headers,
)


REPORT_AT = datetime(2026, 10, 15, 12, 0, tzinfo=timezone(timedelta(hours=-6)))


class B2BCustomersExportTests(unittest.TestCase):
    @classmethod
    def tearDownClass(cls) -> None:
        main.engine.dispose()
        Path(_TEMP_DB.name).unlink(missing_ok=True)
        if _PREVIOUS_DATABASE_URL is None:
            os.environ.pop("DATABASE_URL", None)
        else:
            os.environ["DATABASE_URL"] = _PREVIOUS_DATABASE_URL

    def setUp(self) -> None:
        with main.SessionLocal.begin() as db:
            db.query(B2BARPayment).delete()
            db.query(B2BAROpeningBalance).delete()
            db.query(B2BSalesOrderLine).delete()
            db.query(B2BSalesOrder).delete()
            db.query(B2CSalesOrder).delete()
            db.query(B2BCommercialAgreement).delete()
            db.query(B2BCustomer).delete()
            db.query(Channel).delete()
            db.query(UserSession).delete()
            db.query(UserRole).delete()
            db.query(User).delete()

            self.admin_token = "b2b-customer-export-admin-token"
            self.operator_token = "b2b-customer-export-operator-token"
            admin_user = User(
                username="b2b-export-admin",
                full_name="B2B Export Admin",
                password_hash="not-used-in-test",
                is_active=True,
            )
            operator_user = User(
                username="b2b-export-operator",
                full_name="B2B Export Operator",
                password_hash="not-used-in-test",
                is_active=True,
            )
            db.add_all((admin_user, operator_user))
            db.flush()
            admin_role = db.query(Role).filter(Role.code == "admin").one()
            operator_role = db.query(Role).filter(Role.code == "general_operator").one()
            db.add_all(
                (
                    UserRole(user_id=admin_user.id, role_id=admin_role.id),
                    UserRole(user_id=operator_user.id, role_id=operator_role.id),
                    self._session(admin_user.id, self.admin_token),
                    self._session(operator_user.id, self.operator_token),
                )
            )

            channel = Channel(name="Tiendas Naturales", active=True, applies_to_b2b=True)
            db.add(channel)
            db.flush()

            signed_then_draft = B2BCustomer(
                customer_name="Cliente 001",
                address="Dirección principal",
                province="San José",
                canton="San José",
                district="Carmen",
                legal_name="Cliente Uno S.A.",
                legal_id="001234567890",
                phone="0011223344",
                credit_days=15,
                default_payment_terms="Crédito comercial aprobado",
                channel_id=channel.id,
                active=True,
            )
            inactive_customer = B2BCustomer(
                customer_name="Cliente inactivo",
                legal_id="000000000001",
                phone="00001111",
                active=False,
            )
            generated_only = B2BCustomer(
                customer_name="Cliente generado",
                address='=HYPERLINK("https://example.invalid", "No ejecutar")',
                active=True,
            )
            db.add_all((signed_then_draft, inactive_customer, generated_only))
            db.flush()
            self.customer_id = signed_then_draft.id
            self.inactive_customer_id = inactive_customer.id
            self.generated_customer_id = generated_only.id

            db.add_all(
                (
                    self._agreement(
                        signed_then_draft.id,
                        version=1,
                        status="signed",
                        generated_at=datetime(2026, 9, 1, 11, 30),
                        signed_at=datetime(2026, 9, 3, 9, 45),
                        signature_date=date(2026, 9, 2),
                    ),
                    self._agreement(
                        signed_then_draft.id,
                        version=2,
                        status="draft",
                    ),
                    self._agreement(
                        generated_only.id,
                        version=0,
                        status="generated",
                        generated_at=datetime(2026, 10, 1, 8, 0),
                        signature_date=date(2026, 10, 5),
                        generated_pdf_bytes=b"%PDF-not-a-signature",
                    ),
                )
            )

    @staticmethod
    def _session(user_id: int, raw_token: str) -> UserSession:
        now = datetime.utcnow()
        return UserSession(
            user_id=user_id,
            session_token_hash=hashlib.sha256(raw_token.encode("utf-8")).hexdigest(),
            created_at=now,
            last_seen_at=now,
            expires_at=now + timedelta(hours=1),
        )

    @staticmethod
    def _agreement(
        customer_id: int,
        *,
        version: int,
        status: str,
        generated_at: datetime | None = None,
        signed_at: datetime | None = None,
        signature_date: date | None = None,
        generated_pdf_bytes: bytes | None = None,
    ) -> B2BCommercialAgreement:
        return B2BCommercialAgreement(
            customer_id=customer_id,
            version_number=version,
            status=status,
            template_revision="test",
            generated_at=generated_at,
            signed_at=signed_at,
            signature_date=signature_date,
            generated_pdf_bytes=generated_pdf_bytes,
        )

    @staticmethod
    def _request(permission_codes: set[str]):
        request = mock.Mock()
        request.state.current_user = mock.Mock(id=100, username="export-test")
        request.state.current_permissions = permission_codes
        return request

    @staticmethod
    def _order(
        customer_id: int,
        order_number: str,
        *,
        total: str,
        status: str = "invoiced",
        invoice_date: date | None = None,
        invoiced_at: datetime | None = None,
        delivery_date: date = date(2026, 1, 1),
        observations: str | None = None,
        created_at: datetime = datetime(2026, 1, 1, 8, 0),
    ) -> B2BSalesOrder:
        return B2BSalesOrder(
            order_number=order_number,
            customer_id=customer_id,
            customer_name_snapshot="Snapshot",
            delivery_date=delivery_date,
            status=status,
            total_amount=Decimal(total),
            invoice_date=invoice_date,
            invoiced_at=invoiced_at,
            observations=observations,
            created_at=created_at,
        )

    def _add_ar_order(
        self,
        db,
        order_number: str,
        *,
        total: str,
        invoice_date: date,
        customer_id: int | None = None,
        status: str = "invoiced",
        invoiced_at: datetime | None = datetime(2026, 1, 1, 8, 0),
        opening_balance: str | None = None,
        payments: tuple[str, ...] = (),
    ) -> B2BSalesOrder:
        order = self._order(
            customer_id or self.customer_id,
            order_number,
            total=total,
            status=status,
            invoice_date=invoice_date,
            invoiced_at=invoiced_at,
        )
        db.add(order)
        db.flush()
        if opening_balance is not None:
            db.add(
                B2BAROpeningBalance(
                    sales_order_id=order.id,
                    outstanding_amount=Decimal(opening_balance),
                )
            )
        db.add_all(
            B2BARPayment(
                sales_order_id=order.id,
                payment_date=date(2026, 10, 1),
                amount=Decimal(payment),
            )
            for payment in payments
        )
        return order

    def _export(
        self,
        permissions: set[str] | None = None,
        *,
        generated_at: datetime = REPORT_AT,
    ):
        with mock.patch.object(
            main,
            "current_b2b_customer_export_datetime",
            return_value=generated_at,
        ):
            with main.SessionLocal() as db:
                response = main.export_b2b_customers_xlsx(
                    request=self._request(permissions or {"sales.view"}),
                    db=db,
                )
                self.assertFalse(db.dirty)
                self.assertFalse(db.new)
                self.assertFalse(db.deleted)
                return response

    @staticmethod
    def _http_get(path: str, token: str) -> tuple[int, dict[str, str], bytes]:
        messages: list[dict] = []
        request_sent = False

        async def receive():
            nonlocal request_sent
            if not request_sent:
                request_sent = True
                return {"type": "http.request", "body": b"", "more_body": False}
            return {"type": "http.disconnect"}

        async def send(message):
            messages.append(message)

        scope = {
            "type": "http",
            "asgi": {"version": "3.0", "spec_version": "2.3"},
            "http_version": "1.1",
            "method": "GET",
            "scheme": "http",
            "path": path,
            "raw_path": path.encode("ascii"),
            "query_string": b"",
            "root_path": "",
            "headers": [
                (b"host", b"testserver"),
                (b"cookie", f"erp_session={token}".encode("ascii")),
            ],
            "client": ("127.0.0.1", 12345),
            "server": ("testserver", 80),
        }
        asyncio.run(main.app(scope, receive, send))
        response_start = next(message for message in messages if message["type"] == "http.response.start")
        response_headers = {
            key.decode("latin-1"): value.decode("latin-1")
            for key, value in response_start["headers"]
        }
        response_body = b"".join(
            message.get("body", b"")
            for message in messages
            if message["type"] == "http.response.body"
        )
        return response_start["status"], response_headers, response_body

    def test_export_contains_exact_columns_and_one_row_per_active_or_inactive_customer(self) -> None:
        response = self._export()
        workbook = load_workbook(io.BytesIO(response.body), data_only=False)
        worksheet = workbook["Clientes B2B"]

        self.assertEqual(
            tuple(cell.value for cell in worksheet[1]),
            b2b_customer_export_headers(2026),
        )
        self.assertEqual(tuple(cell.value for cell in worksheet[1][:14]), B2B_CUSTOMER_EXPORT_HEADERS)
        self.assertEqual(worksheet.max_column, 20)
        self.assertEqual(
            tuple(cell.value for cell in worksheet[1][17:20]),
            (
                "Cantidad de facturas vencidas",
                "Monto total vencido",
                "Días de mayor atraso",
            ),
        )
        self.assertEqual(worksheet.max_row, 4)
        removed_headers = {
            "Versión del último acuerdo",
            "Fecha de generación del último acuerdo",
            "Fecha de firma registrada",
        }
        self.assertTrue(removed_headers.isdisjoint(cell.value for cell in worksheet[1]))
        rows_by_name = {
            worksheet.cell(row=row_index, column=1).value: row_index
            for row_index in range(2, worksheet.max_row + 1)
        }
        self.assertEqual(
            set(rows_by_name),
            {"Cliente 001", "Cliente generado", "Cliente inactivo"},
        )
        self.assertEqual(worksheet.cell(rows_by_name["Cliente 001"], 3).value, "Activo")
        self.assertEqual(worksheet.cell(rows_by_name["Cliente inactivo"], 3).value, "Inactivo")

    def test_reordered_customer_values_align_with_the_requested_headers(self) -> None:
        worksheet = load_workbook(io.BytesIO(self._export().body))["Clientes B2B"]
        customer_row = self._row_for_customer(worksheet, "Cliente 001")
        self.assertEqual(
            tuple(worksheet.cell(customer_row, column).value for column in range(1, 15)),
            (
                "Cliente 001",
                "Tiendas Naturales",
                "Activo",
                "Cliente Uno S.A.",
                "001234567890",
                "Dirección principal",
                "San José",
                "San José",
                "Carmen",
                "0011223344",
                "Crédito comercial aprobado",
                15,
                "draft",
                "Sí",
            ),
        )

    def test_latest_agreement_and_any_registered_signature_are_independent(self) -> None:
        worksheet = load_workbook(io.BytesIO(self._export().body))["Clientes B2B"]
        rows = {
            worksheet.cell(row=row_index, column=1).value: row_index
            for row_index in range(2, worksheet.max_row + 1)
        }
        signed_customer_row = rows["Cliente 001"]
        generated_customer_row = rows["Cliente generado"]

        self.assertEqual(worksheet.cell(signed_customer_row, 13).value, "draft")
        self.assertEqual(worksheet.cell(signed_customer_row, 14).value, "Sí")

        self.assertEqual(worksheet.cell(generated_customer_row, 13).value, "generated")
        self.assertEqual(worksheet.cell(generated_customer_row, 14).value, NO_REGISTERED_SIGNATURE)

    def test_identifiers_are_text_and_empty_values_remain_empty(self) -> None:
        worksheet = load_workbook(io.BytesIO(self._export().body))["Clientes B2B"]
        rows = {
            worksheet.cell(row=row_index, column=1).value: row_index
            for row_index in range(2, worksheet.max_row + 1)
        }
        populated_row = rows["Cliente 001"]
        empty_row = rows["Cliente inactivo"]

        self.assertEqual(worksheet.cell(populated_row, 5).value, "001234567890")
        self.assertEqual(worksheet.cell(populated_row, 10).value, "0011223344")
        self.assertEqual(worksheet.cell(populated_row, 5).number_format, "@")
        self.assertEqual(worksheet.cell(populated_row, 10).number_format, "@")
        self.assertIsNone(worksheet.cell(empty_row, 2).value)
        self.assertIsNone(worksheet.cell(empty_row, 6).value)
        self.assertIsNone(worksheet.cell(empty_row, 11).value)
        self.assertIsNone(worksheet.cell(empty_row, 13).value)
        self.assertEqual(worksheet.cell(empty_row, 14).value, NO_REGISTERED_SIGNATURE)

    def test_user_controlled_text_cannot_become_an_excel_formula(self) -> None:
        worksheet = load_workbook(io.BytesIO(self._export().body), data_only=False)["Clientes B2B"]
        generated_row = next(
            row_index
            for row_index in range(2, worksheet.max_row + 1)
            if worksheet.cell(row=row_index, column=1).value == "Cliente generado"
        )
        address_cell = worksheet.cell(row=generated_row, column=6)

        self.assertEqual(
            address_cell.value,
            '=HYPERLINK("https://example.invalid", "No ejecutar")',
        )
        self.assertEqual(address_cell.data_type, "s")

    def test_workbook_has_filter_frozen_header_and_download_headers(self) -> None:
        response = self._export()
        worksheet = load_workbook(io.BytesIO(response.body))["Clientes B2B"]

        self.assertEqual(worksheet.freeze_panes, "A2")
        self.assertEqual(worksheet.auto_filter.ref, "A1:T4")
        self.assertTrue(worksheet["A1"].font.bold)
        self.assertEqual(
            response.media_type,
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        self.assertEqual(
            response.headers["content-disposition"],
            'attachment; filename="Reporte_Clientes_B2B_2026-10-15.xlsx"',
        )

    def test_export_requires_sales_view(self) -> None:
        with self.assertRaises(HTTPException) as raised:
            self._export({"product.view"})
        self.assertEqual(raised.exception.status_code, 403)

    def test_real_http_route_returns_valid_excel_and_rejects_missing_permission(self) -> None:
        with mock.patch.object(
            main,
            "current_b2b_customer_export_datetime",
            return_value=REPORT_AT,
        ):
            status_code, headers, body = self._http_get(
                "/b2b/customers/export.xlsx",
                self.admin_token,
            )
        self.assertEqual(status_code, 200)
        self.assertEqual(
            headers["content-type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        self.assertEqual(
            headers["content-disposition"],
            'attachment; filename="Reporte_Clientes_B2B_2026-10-15.xlsx"',
        )
        workbook = load_workbook(io.BytesIO(body))
        self.assertEqual(workbook["Clientes B2B"].max_column, 20)

        forbidden_status, _forbidden_headers, forbidden_body = self._http_get(
            "/b2b/customers/export.xlsx",
            self.operator_token,
        )
        self.assertEqual(forbidden_status, 403)
        self.assertIn(b"You do not have access to this action", forbidden_body)

    def test_dynamic_year_headers_and_filename_use_costa_rica_time(self) -> None:
        still_2026_in_costa_rica = datetime(2027, 1, 1, 4, 30, tzinfo=timezone.utc)
        response_2026 = self._export(generated_at=still_2026_in_costa_rica)
        worksheet_2026 = load_workbook(io.BytesIO(response_2026.body))["Clientes B2B"]
        self.assertEqual(worksheet_2026.cell(1, 15).value, "Ventas 2025")
        self.assertEqual(worksheet_2026.cell(1, 16).value, "Ventas acumuladas 2026")
        self.assertEqual(
            response_2026.headers["content-disposition"],
            'attachment; filename="Reporte_Clientes_B2B_2026-12-31.xlsx"',
        )

        already_2027_in_costa_rica = datetime(2027, 1, 1, 6, 30, tzinfo=timezone.utc)
        response_2027 = self._export(generated_at=already_2027_in_costa_rica)
        worksheet_2027 = load_workbook(io.BytesIO(response_2027.body))["Clientes B2B"]
        self.assertEqual(worksheet_2027.cell(1, 15).value, "Ventas 2026")
        self.assertEqual(worksheet_2027.cell(1, 16).value, "Ventas acumuladas 2027")

    def test_sales_use_canonical_invoice_date_precedence(self) -> None:
        with main.SessionLocal.begin() as db:
            db.add_all(
                (
                    self._order(
                        self.customer_id,
                        "DATE-INV-DATE",
                        total="100",
                        invoice_date=date(2025, 5, 1),
                        invoiced_at=datetime(2026, 5, 1, 8, 0),
                        delivery_date=date(2026, 5, 1),
                        created_at=datetime(2026, 5, 1, 8, 0),
                    ),
                    self._order(
                        self.customer_id,
                        "DATE-INVOICED-AT",
                        total="200",
                        invoiced_at=datetime(2025, 6, 1, 8, 0),
                        delivery_date=date(2026, 6, 1),
                        observations="Historical CSV import 123",
                        created_at=datetime(2026, 6, 1, 8, 0),
                    ),
                    self._order(
                        self.customer_id,
                        "DATE-HISTORICAL",
                        total="300",
                        delivery_date=date(2026, 7, 1),
                        observations="Historical CSV import 456",
                        created_at=datetime(2025, 7, 1, 8, 0),
                    ),
                    self._order(
                        self.customer_id,
                        "DATE-CREATED",
                        total="400",
                        delivery_date=date(2026, 8, 1),
                        observations="Normal order",
                        created_at=datetime(2025, 8, 1, 8, 0),
                    ),
                )
            )

        worksheet = load_workbook(io.BytesIO(self._export().body))["Clientes B2B"]
        customer_row = self._row_for_customer(worksheet, "Cliente 001")
        self.assertEqual(worksheet.cell(customer_row, 15).value, 700)
        self.assertEqual(worksheet.cell(customer_row, 16).value, 300)
        self.assertEqual(
            self._cell_date(worksheet.cell(customer_row, 17).value),
            date(2026, 7, 1),
        )

    def test_only_invoiced_orders_and_no_future_current_year_sales_are_counted(self) -> None:
        with main.SessionLocal.begin() as db:
            db.add_all(
                (
                    self._order(
                        self.generated_customer_id,
                        "STATUS-INVOICED",
                        total="50",
                        invoice_date=date(2026, 10, 15),
                    ),
                    self._order(
                        self.generated_customer_id,
                        "STATUS-DRAFT",
                        total="100",
                        status="draft",
                        invoice_date=date(2026, 10, 15),
                    ),
                    self._order(
                        self.generated_customer_id,
                        "STATUS-IN-PROCESS",
                        total="200",
                        status="in_process",
                        invoice_date=date(2026, 10, 15),
                    ),
                    self._order(
                        self.generated_customer_id,
                        "STATUS-CANCELLED",
                        total="300",
                        status="cancelled",
                        invoice_date=date(2026, 10, 15),
                    ),
                    self._order(
                        self.generated_customer_id,
                        "STATUS-FUTURE",
                        total="400",
                        invoice_date=date(2026, 10, 16),
                    ),
                )
            )

        worksheet = load_workbook(io.BytesIO(self._export().body))["Clientes B2B"]
        customer_row = self._row_for_customer(worksheet, "Cliente generado")
        self.assertEqual(worksheet.cell(customer_row, 15).value, 0)
        self.assertEqual(worksheet.cell(customer_row, 16).value, 50)

    def test_previous_year_uses_complete_calendar_year_and_current_year_is_accumulated(self) -> None:
        with main.SessionLocal.begin() as db:
            db.add_all(
                (
                    self._order(
                        self.customer_id,
                        "BOUNDARY-BEFORE",
                        total="10",
                        invoice_date=date(2024, 12, 31),
                    ),
                    self._order(
                        self.customer_id,
                        "BOUNDARY-PREVIOUS-START",
                        total="20",
                        invoice_date=date(2025, 1, 1),
                    ),
                    self._order(
                        self.customer_id,
                        "BOUNDARY-PREVIOUS-END",
                        total="30",
                        invoice_date=date(2025, 12, 31),
                    ),
                    self._order(
                        self.customer_id,
                        "BOUNDARY-CURRENT-START",
                        total="40",
                        invoice_date=date(2026, 1, 1),
                    ),
                    self._order(
                        self.customer_id,
                        "BOUNDARY-CURRENT-TODAY",
                        total="50",
                        invoice_date=date(2026, 10, 15),
                    ),
                )
            )

        worksheet = load_workbook(io.BytesIO(self._export().body))["Clientes B2B"]
        customer_row = self._row_for_customer(worksheet, "Cliente 001")
        self.assertEqual(worksheet.cell(customer_row, 15).value, 50)
        self.assertEqual(worksheet.cell(customer_row, 16).value, 90)

    def test_order_total_is_counted_once_without_b2c_lines_or_payments(self) -> None:
        with main.SessionLocal.begin() as db:
            order = self._order(
                self.customer_id,
                "NO-DUPLICATION",
                total="150",
                invoice_date=date(2026, 4, 1),
            )
            db.add(order)
            db.flush()
            db.add_all(
                (
                    B2BSalesOrderLine(
                        sales_order_id=order.id,
                        line_number=1,
                        sku_snapshot="SKU-1",
                        description_snapshot="Línea 1",
                        unit_price_snapshot=Decimal("50"),
                        quantity=Decimal("1"),
                        line_total=Decimal("50"),
                    ),
                    B2BSalesOrderLine(
                        sales_order_id=order.id,
                        line_number=2,
                        sku_snapshot="SKU-2",
                        description_snapshot="Línea 2",
                        unit_price_snapshot=Decimal("100"),
                        quantity=Decimal("1"),
                        line_total=Decimal("100"),
                    ),
                    B2BARPayment(
                        sales_order_id=order.id,
                        payment_date=date(2026, 4, 15),
                        amount=Decimal("50"),
                    ),
                    B2BARPayment(
                        sales_order_id=order.id,
                        payment_date=date(2026, 5, 15),
                        amount=Decimal("100"),
                    ),
                    B2CSalesOrder(
                        order_number="B2C-NOT-INCLUDED",
                        order_date=date(2026, 4, 1),
                        channel="Retail",
                        status="invoiced",
                        total_amount=Decimal("9999"),
                    ),
                )
            )

        worksheet = load_workbook(io.BytesIO(self._export().body))["Clientes B2B"]
        customer_row = self._row_for_customer(worksheet, "Cliente 001")
        self.assertEqual(worksheet.cell(customer_row, 16).value, 150)

    def test_customer_without_invoiced_sales_has_zero_totals_and_blank_last_date(self) -> None:
        worksheet = load_workbook(io.BytesIO(self._export().body))["Clientes B2B"]
        customer_row = self._row_for_customer(worksheet, "Cliente inactivo")
        self.assertEqual(worksheet.cell(customer_row, 15).value, 0)
        self.assertEqual(worksheet.cell(customer_row, 16).value, 0)
        self.assertIsNone(worksheet.cell(customer_row, 17).value)
        self.assertEqual(worksheet.cell(customer_row, 18).value, 0)
        self.assertEqual(worksheet.cell(customer_row, 19).value, 0)
        self.assertEqual(worksheet.cell(customer_row, 20).value, 0)

    def test_overdue_metrics_use_pending_balances_and_count_each_invoice_once(self) -> None:
        with main.SessionLocal.begin() as db:
            self._add_ar_order(
                db,
                "OVERDUE-UNPAID",
                total="100000",
                invoice_date=date(2026, 8, 1),
            )
            self._add_ar_order(
                db,
                "OVERDUE-PARTIAL",
                total="80000",
                invoice_date=date(2026, 9, 1),
                payments=("10000", "10000"),
            )

        worksheet = load_workbook(io.BytesIO(self._export().body))["Clientes B2B"]
        customer_row = self._row_for_customer(worksheet, "Cliente 001")
        self.assertEqual(worksheet.cell(customer_row, 18).value, 2)
        self.assertEqual(worksheet.cell(customer_row, 19).value, 160000)
        self.assertEqual(worksheet.cell(customer_row, 20).value, 60)

    def test_paid_cancelled_draft_and_due_today_invoices_are_not_overdue(self) -> None:
        with main.SessionLocal.begin() as db:
            self._add_ar_order(
                db,
                "PAID-IN-FULL",
                total="50000",
                invoice_date=date(2026, 9, 1),
                payments=("50000",),
            )
            self._add_ar_order(
                db,
                "CANCELLED-INVOICE",
                total="100000",
                status="cancelled",
                invoice_date=date(2026, 8, 1),
            )
            self._add_ar_order(
                db,
                "DRAFT-ORDER",
                total="100000",
                status="draft",
                invoice_date=date(2026, 8, 1),
            )
            self._add_ar_order(
                db,
                "DUE-TODAY",
                total="100000",
                invoice_date=date(2026, 9, 30),
            )

        worksheet = load_workbook(io.BytesIO(self._export().body))["Clientes B2B"]
        customer_row = self._row_for_customer(worksheet, "Cliente 001")
        self.assertEqual(worksheet.cell(customer_row, 18).value, 0)
        self.assertEqual(worksheet.cell(customer_row, 19).value, 0)
        self.assertEqual(worksheet.cell(customer_row, 20).value, 0)

    def test_invoice_due_yesterday_has_one_day_overdue(self) -> None:
        with main.SessionLocal.begin() as db:
            self._add_ar_order(
                db,
                "DUE-YESTERDAY",
                total="25000",
                invoice_date=date(2026, 9, 29),
            )

        worksheet = load_workbook(io.BytesIO(self._export().body))["Clientes B2B"]
        customer_row = self._row_for_customer(worksheet, "Cliente 001")
        self.assertEqual(worksheet.cell(customer_row, 18).value, 1)
        self.assertEqual(worksheet.cell(customer_row, 19).value, 25000)
        self.assertEqual(worksheet.cell(customer_row, 20).value, 1)

    def test_historical_opening_balance_and_payments_use_current_ar_balance(self) -> None:
        with main.SessionLocal.begin() as db:
            self._add_ar_order(
                db,
                "HISTORICAL-OPENING-BALANCE",
                total="100000",
                invoice_date=date(2026, 9, 1),
                invoiced_at=None,
                opening_balance="70000",
                payments=("20000",),
            )

        worksheet = load_workbook(io.BytesIO(self._export().body))["Clientes B2B"]
        customer_row = self._row_for_customer(worksheet, "Cliente 001")
        self.assertEqual(worksheet.cell(customer_row, 18).value, 1)
        self.assertEqual(worksheet.cell(customer_row, 19).value, 50000)
        self.assertEqual(worksheet.cell(customer_row, 20).value, 29)

    def test_overdue_amount_and_days_match_accounts_receivable(self) -> None:
        with main.SessionLocal.begin() as db:
            self._add_ar_order(
                db,
                "MATCH-AR",
                total="100.0000",
                invoice_date=date(2026, 9, 1),
                payments=("99.9999",),
            )

        with main.SessionLocal() as db:
            dashboard = build_accounts_receivable_dashboard(db, today=REPORT_AT.date())
            ar_row = next(row for row in dashboard.rows if row.order_number == "MATCH-AR")

        worksheet = load_workbook(io.BytesIO(self._export().body))["Clientes B2B"]
        customer_row = self._row_for_customer(worksheet, "Cliente 001")
        self.assertEqual(worksheet.cell(customer_row, 18).value, 1)
        self.assertEqual(
            Decimal(str(worksheet.cell(customer_row, 19).value)),
            ar_row.pending_amount,
        )
        self.assertEqual(worksheet.cell(customer_row, 20).value, -ar_row.days_remaining)

    def test_overdue_metrics_match_ar_dashboard_for_every_customer(self) -> None:
        with main.SessionLocal.begin() as db:
            self._add_ar_order(
                db,
                "AR-PARITY-PARTIAL",
                total="100000",
                invoice_date=date(2026, 9, 1),
                payments=("25000", "15000"),
            )
            self._add_ar_order(
                db,
                "AR-PARITY-OLD",
                total="50000",
                invoice_date=date(2026, 8, 1),
            )
            self._add_ar_order(
                db,
                "AR-PARITY-SECOND-CUSTOMER",
                total="70000",
                invoice_date=date(2026, 10, 14),
                customer_id=self.generated_customer_id,
            )

        with main.SessionLocal() as db:
            dashboard = build_accounts_receivable_dashboard(db, today=REPORT_AT.date())
            expected_by_customer_id: dict[int, tuple[int, Decimal, int]] = {}
            for row in dashboard.rows:
                if row.status_code != "overdue":
                    continue
                count, amount, max_days = expected_by_customer_id.get(
                    row.customer_id,
                    (0, Decimal("0"), 0),
                )
                expected_by_customer_id[row.customer_id] = (
                    count + 1,
                    amount + row.pending_amount,
                    max(max_days, -(row.days_remaining or 0)),
                )

        worksheet = load_workbook(io.BytesIO(self._export().body))["Clientes B2B"]
        customer_ids_by_name = {
            "Cliente 001": self.customer_id,
            "Cliente generado": self.generated_customer_id,
            "Cliente inactivo": self.inactive_customer_id,
        }
        for customer_name, customer_id in customer_ids_by_name.items():
            row_index = self._row_for_customer(worksheet, customer_name)
            expected_count, expected_amount, expected_days = expected_by_customer_id.get(
                customer_id,
                (0, Decimal("0"), 0),
            )
            self.assertEqual(worksheet.cell(row_index, 18).value, expected_count)
            self.assertEqual(
                Decimal(str(worksheet.cell(row_index, 19).value)),
                expected_amount,
            )
            self.assertEqual(worksheet.cell(row_index, 20).value, expected_days)

    def test_overdue_reference_date_uses_costa_rica_calendar_date(self) -> None:
        with main.SessionLocal.begin() as db:
            self._add_ar_order(
                db,
                "COSTA-RICA-DATE",
                total="1000",
                invoice_date=date(2026, 12, 15),
                invoiced_at=datetime(2026, 12, 15, 8, 0),
            )

        still_december_31_in_costa_rica = datetime(2027, 1, 1, 5, 30, tzinfo=timezone.utc)
        worksheet = load_workbook(
            io.BytesIO(self._export(generated_at=still_december_31_in_costa_rica).body)
        )["Clientes B2B"]
        customer_row = self._row_for_customer(worksheet, "Cliente 001")
        self.assertEqual(worksheet.cell(customer_row, 20).value, 1)

        january_1_in_costa_rica = datetime(2027, 1, 1, 6, 30, tzinfo=timezone.utc)
        worksheet = load_workbook(
            io.BytesIO(self._export(generated_at=january_1_in_costa_rica).body)
        )["Clientes B2B"]
        customer_row = self._row_for_customer(worksheet, "Cliente 001")
        self.assertEqual(worksheet.cell(customer_row, 20).value, 2)

    def test_last_invoice_date_uses_full_history_and_excludes_cancelled_orders(self) -> None:
        with main.SessionLocal.begin() as db:
            db.add_all(
                (
                    self._order(
                        self.generated_customer_id,
                        "LAST-HISTORICAL-INVOICED",
                        total="25",
                        invoice_date=date(2020, 2, 3),
                    ),
                    self._order(
                        self.generated_customer_id,
                        "LAST-NEWER-CANCELLED",
                        total="999",
                        status="cancelled",
                        invoice_date=date(2026, 9, 30),
                    ),
                )
            )

        worksheet = load_workbook(io.BytesIO(self._export().body))["Clientes B2B"]
        customer_row = self._row_for_customer(worksheet, "Cliente generado")
        self.assertEqual(worksheet.cell(customer_row, 15).value, 0)
        self.assertEqual(worksheet.cell(customer_row, 16).value, 0)
        self.assertEqual(
            self._cell_date(worksheet.cell(customer_row, 17).value),
            date(2020, 2, 3),
        )

    def test_sales_and_last_invoice_cells_have_numeric_and_date_formats(self) -> None:
        with main.SessionLocal.begin() as db:
            db.add(
                self._order(
                    self.customer_id,
                    "FORMATTING",
                    total="1234.50",
                    invoice_date=date(2026, 3, 4),
                )
            )

        worksheet = load_workbook(io.BytesIO(self._export().body))["Clientes B2B"]
        customer_row = self._row_for_customer(worksheet, "Cliente 001")
        sales_cell = worksheet.cell(customer_row, 16)
        last_invoice_cell = worksheet.cell(customer_row, 17)
        self.assertEqual(sales_cell.data_type, "n")
        self.assertEqual(sales_cell.number_format, "₡ #,##0.00")
        self.assertEqual(last_invoice_cell.number_format, "dd/mm/yyyy")
        self.assertEqual(self._cell_date(last_invoice_cell.value), date(2026, 3, 4))

    def test_overdue_cells_have_numeric_types_and_currency_format(self) -> None:
        with main.SessionLocal.begin() as db:
            self._add_ar_order(
                db,
                "OVERDUE-FORMATTING",
                total="1234.50",
                invoice_date=date(2026, 9, 1),
            )

        worksheet = load_workbook(io.BytesIO(self._export().body))["Clientes B2B"]
        customer_row = self._row_for_customer(worksheet, "Cliente 001")
        count_cell = worksheet.cell(customer_row, 18)
        amount_cell = worksheet.cell(customer_row, 19)
        days_cell = worksheet.cell(customer_row, 20)
        self.assertEqual(count_cell.data_type, "n")
        self.assertEqual(amount_cell.data_type, "n")
        self.assertEqual(amount_cell.number_format, "₡ #,##0.00")
        self.assertEqual(days_cell.data_type, "n")

    def test_sales_and_ar_are_loaded_without_one_query_per_customer(self) -> None:
        with main.SessionLocal.begin() as db:
            db.add_all(
                B2BCustomer(customer_name=f"Cliente rendimiento {index:02d}", active=True)
                for index in range(12)
            )

        all_selects: list[str] = []
        sales_selects: list[str] = []

        def capture_statement(_connection, _cursor, statement, _parameters, _context, _executemany):
            normalized = " ".join(statement.lower().split())
            if normalized.startswith("select"):
                all_selects.append(normalized)
            if " from b2b_sales_orders" in normalized:
                sales_selects.append(normalized)

        event.listen(main.engine, "before_cursor_execute", capture_statement)
        try:
            self._export()
        finally:
            event.remove(main.engine, "before_cursor_execute", capture_statement)

        self.assertEqual(len(sales_selects), 2)
        self.assertEqual(len(all_selects), 5)
        self.assertTrue(
            any("group by b2b_sales_orders.customer_id" in statement for statement in sales_selects)
        )

    @staticmethod
    def _row_for_customer(worksheet, customer_name: str) -> int:
        return next(
            row_index
            for row_index in range(2, worksheet.max_row + 1)
            if worksheet.cell(row=row_index, column=1).value == customer_name
        )

    @staticmethod
    def _cell_date(value) -> date | None:
        if isinstance(value, datetime):
            return value.date()
        return value

    def test_customer_list_shows_export_button(self) -> None:
        template = Path("app/templates/b2b_customers_list.html").read_text(encoding="utf-8")
        self.assertIn("Exportar clientes B2B", template)
        self.assertIn('href="/b2b/customers/export.xlsx"', template)
        self.assertIn('can(request, "sales.view")', template)


if __name__ == "__main__":
    unittest.main()
