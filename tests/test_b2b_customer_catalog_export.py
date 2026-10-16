import csv
import io
import os
import tempfile
import unittest
from datetime import timezone
from decimal import Decimal
from pathlib import Path
from unittest import mock

from fastapi import HTTPException
from openpyxl import load_workbook


_PREVIOUS_DATABASE_URL = os.environ.get("DATABASE_URL")
_TEMP_DB = tempfile.NamedTemporaryFile(prefix="b2b_catalog_export_", suffix=".db", delete=False)
_TEMP_DB.close()
os.environ["DATABASE_URL"] = f"sqlite:///{Path(_TEMP_DB.name).as_posix()}"

with mock.patch("zoneinfo.ZoneInfo", return_value=timezone.utc):
    from app import main  # noqa: E402
from app.models import (  # noqa: E402
    AppSettings,
    B2BCustomer,
    B2BCustomerProduct,
    Channel,
    Permission,
    Product,
    ProductCategory,
    Role,
    RolePermission,
    Route,
)
from app.services.b2b_catalog_export_service import (  # noqa: E402
    format_crc_catalog_price,
    safe_b2b_catalog_pdf_filename,
)
from app.services.b2b_commercial_inactivity_service import (  # noqa: E402
    AUTHORIZED_B2B_CHANNEL_NAMES,
)
from app.services.b2b_updated_catalog_export_service import (  # noqa: E402
    EXCLUDED_SERVICE_SKUS,
    MISSING_B2C_PRICE_LABEL,
    PRODUCT_SCOPE_ALL,
    PRODUCT_SCOPE_BULK,
    PRODUCT_SCOPE_PACKAGED,
    REVIEW_HEADERS,
    UPDATED_B2B_CATALOG_HEADERS,
)


class B2BCustomerCatalogExportTests(unittest.TestCase):
    @classmethod
    def tearDownClass(cls) -> None:
        main.engine.dispose()
        Path(_TEMP_DB.name).unlink(missing_ok=True)
        if _PREVIOUS_DATABASE_URL is None:
            os.environ.pop("DATABASE_URL", None)
        else:
            os.environ["DATABASE_URL"] = _PREVIOUS_DATABASE_URL

    def setUp(self) -> None:
        with main.SessionLocal.begin() as db:
            db.query(B2BCustomerProduct).delete()
            db.query(B2BCustomer).delete()
            db.query(Channel).delete()
            db.query(Product).delete()
            db.query(Route).delete()
            db.query(ProductCategory).delete()
            db.query(AppSettings).delete()
            channel = Channel(name="Tiendas Naturales (Minisuper)", active=True, applies_to_b2b=True)
            base_packaging_route = Route(
                code="R-BASE-PACKAGING",
                name="Ruta base de envasado",
                process_type="packaging",
                active=True,
            )
            db.add_all((channel, base_packaging_route))
            db.flush()
            self.base_packaging_route_id = base_packaging_route.id
            customer_a = B2BCustomer(
                customer_name="Café Cliente A",
                legal_name="Grupo Café Cliente A S.A.",
                legal_id="3-101-123456",
                phone="24797018",
                address="Costado Norte del Parque, La Fortuna\nAlajuela",
                province="Alajuela",
                canton="San Carlos",
                district="La Fortuna",
                credit_days=10,
                channel_id=channel.id,
                active=True,
            )
            customer_b = B2BCustomer(customer_name="Cliente B")
            db.add_all((customer_a, customer_b))
            db.flush()
            self.customer_a_id = customer_a.id
            self.customer_b_id = customer_b.id
            db.add_all(
                (
                    B2BCustomerProduct(
                        customer_id=customer_a.id,
                        sku="A-001",
                        description="Té limón",
                        distributor_price=Decimal("1250.5000"),
                        active=True,
                    ),
                    B2BCustomerProduct(
                        customer_id=customer_a.id,
                        sku="A-002",
                        description="Producto inactivo",
                        distributor_price=Decimal("800.0000"),
                        active=False,
                    ),
                    B2BCustomerProduct(
                        customer_id=customer_b.id,
                        sku="B-SECRET",
                        description="Otro cliente",
                        distributor_price=Decimal("9999.0000"),
                        active=True,
                    ),
                )
            )
            db.add(
                Product(
                    sku="A-001",
                    name="Té limón",
                    b2c_price=Decimal("3200.0000"),
                    default_route_id=base_packaging_route.id,
                    is_manufactured=True,
                )
            )

    def _export_customer(self, customer_id: int):
        with main.SessionLocal() as db, mock.patch.object(main, "require_permission") as permission:
            response = main.export_b2b_customer_products_csv(
                customer_id=customer_id,
                request=mock.Mock(),
                db=db,
            )
            permission.assert_called_once_with(mock.ANY, "sales.view")
            self.assertFalse(db.dirty)
            self.assertFalse(db.new)
            self.assertFalse(db.deleted)
            return response

    def _export_updated_catalog(
        self,
        discount_percent: str = "26.8",
        permission_codes: set[str] | None = None,
        product_scope: str = PRODUCT_SCOPE_ALL,
    ):
        request = self._request_with_permissions(permission_codes or {"sales.view"})
        with main.SessionLocal() as db:
            response = main.export_updated_b2b_catalog_xlsx(
                request=request,
                discount_percent=discount_percent,
                product_scope=product_scope,
                db=db,
            )
            self.assertFalse(db.dirty)
            self.assertFalse(db.new)
            self.assertFalse(db.deleted)
            return response

    def _role_permissions(self, role_code: str) -> set[str]:
        with main.SessionLocal() as db:
            return {
                code
                for (code,) in (
                    db.query(Permission.code)
                    .join(RolePermission, RolePermission.permission_id == Permission.id)
                    .join(Role, Role.id == RolePermission.role_id)
                    .filter(Role.code == role_code, Role.active.is_(True), Permission.active.is_(True))
                    .all()
                )
            }

    def _export_customer_pdf(self, customer_id: int, permission_codes: set[str]):
        request = self._request_with_permissions(permission_codes)
        with main.SessionLocal() as db:
            response = main.export_b2b_customer_products_pdf(
                customer_id=customer_id,
                request=request,
                db=db,
            )
            self.assertFalse(db.dirty)
            self.assertFalse(db.new)
            self.assertFalse(db.deleted)
            return response

    @staticmethod
    def _request_with_permissions(permission_codes: set[str]):
        request = mock.Mock()
        request.state.current_user = mock.Mock(id=123, username="catalog-test-user")
        request.state.current_permissions = permission_codes
        return request

    def test_export_is_utf8_bom_and_strictly_scoped_to_customer(self) -> None:
        response = self._export_customer(self.customer_a_id)
        self.assertTrue(response.body.startswith(b"\xef\xbb\xbf"))
        self.assertEqual(
            response.headers["content-disposition"],
            f'attachment; filename="b2b_customer_{self.customer_a_id}_catalog.csv"',
        )
        rows = list(csv.reader(io.StringIO(response.body.decode("utf-8-sig"))))
        self.assertEqual(rows[0], ["Catalogo:", "Café Cliente A"])
        self.assertEqual(rows[1], [])
        self.assertNotIn(["Campo", "Valor", "Campo", "Valor"], rows)
        self.assertEqual(rows[2], ["Customer name", "Café Cliente A", "Legal name", "Grupo Café Cliente A S.A."])
        self.assertEqual(rows[3], ["Legal ID", "3-101-123456", "Phone", "24797018"])
        self.assertEqual(
            rows[4],
            [
                "Address",
                "Costado Norte del Parque, La Fortuna\nAlajuela",
                "B2B Channel",
                "Tiendas Naturales (Minisuper)",
            ],
        )
        self.assertEqual(rows[5], ["Province", "Alajuela", "Credit days", "10"])
        self.assertEqual(rows[6], ["Canton", "San Carlos", "Active", "yes"])
        self.assertEqual(rows[7], ["District", "La Fortuna", "", ""])
        self.assertEqual(rows[8], [])
        self.assertEqual(rows[9], ["Codigo", "Descripcion", "Precio", "Estado"])
        self.assertEqual(
            rows[10:],
            [
                ["A-001", "Té limón", "1250.5000", "Activo"],
                ["A-002", "Producto inactivo", "800.0000", "Inactivo"],
            ],
        )
        self.assertIn('"Costado Norte del Parque, La Fortuna', response.body.decode("utf-8-sig"))
        self.assertNotIn("B-SECRET", response.body.decode("utf-8-sig"))

    def test_empty_customer_catalog_exports_only_headers(self) -> None:
        with main.SessionLocal.begin() as db:
            empty_customer = B2BCustomer(customer_name="Sin catálogo")
            db.add(empty_customer)
            db.flush()
            empty_customer_id = empty_customer.id

        response = self._export_customer(empty_customer_id)
        rows = list(csv.reader(io.StringIO(response.body.decode("utf-8-sig"))))
        self.assertEqual(rows[0], ["Catalogo:", "Sin catálogo"])
        self.assertEqual(rows[4], ["Address", "", "B2B Channel", ""])
        self.assertEqual(rows[8], [])
        self.assertEqual(rows[9], ["Codigo", "Descripcion", "Precio", "Estado"])
        self.assertEqual(len(rows), 10)

    def test_templates_show_catalog_and_download_labels(self) -> None:
        customer_list = Path("app/templates/b2b_customers_list.html").read_text(encoding="utf-8")
        catalog = Path("app/templates/b2b_customer_products.html").read_text(encoding="utf-8")
        self.assertIn(">Catálogo</a>", customer_list)
        self.assertNotIn(">Products</a>", customer_list)
        self.assertIn("Descargar CSV", catalog)
        self.assertIn("Descargar PDF", catalog)
        self.assertIn("products/export.csv", catalog)
        self.assertIn("products/export.pdf", catalog)
        self.assertEqual(catalog.count("products/export.pdf"), 1)

    def test_updated_catalog_button_uses_existing_page_and_discount_form(self) -> None:
        customer_list = Path("app/templates/b2b_customers_list.html").read_text(encoding="utf-8")
        self.assertIn("Catálogo B2B actualizado", customer_list)
        self.assertIn('action="/b2b/customers/catalog-updated.xlsx"', customer_list)
        self.assertIn('name="discount_percent"', customer_list)
        self.assertIn('name="product_scope"', customer_list)
        self.assertIn('value="packaged"', customer_list)
        self.assertIn('value="bulk"', customer_list)
        self.assertIn('value="all"', customer_list)
        self.assertIn("Envasados", customer_list)
        self.assertIn("A granel", customer_list)
        self.assertIn("Todos los productos comerciales", customer_list)
        self.assertIn('min="0"', customer_list)
        self.assertIn('max="100"', customer_list)
        self.assertIn("Descargar Excel", customer_list)

    def test_updated_catalog_export_is_unique_scoped_calculated_and_read_only(self) -> None:
        with main.SessionLocal.begin() as db:
            category = ProductCategory(name="Infusiones", active=True)
            second_authorized_channel = Channel(
                name=AUTHORIZED_B2B_CHANNEL_NAMES[1],
                active=True,
                applies_to_b2b=True,
            )
            other_channel = Channel(
                name="Canal fuera del alcance",
                active=True,
                applies_to_b2b=True,
            )
            db.add_all((category, second_authorized_channel, other_channel))
            db.flush()

            canonical_product = db.query(Product).filter(Product.sku == "A-001").one()
            canonical_product.category_id = category.id
            duplicate_customer = B2BCustomer(
                customer_name="Cliente autorizado duplicado",
                channel_id=second_authorized_channel.id,
                active=True,
            )
            inactive_customer = B2BCustomer(
                customer_name="Cliente autorizado inactivo",
                channel_id=second_authorized_channel.id,
                active=False,
            )
            other_channel_customer = B2BCustomer(
                customer_name="Cliente de otro canal",
                channel_id=other_channel.id,
                active=True,
            )
            db.add_all((duplicate_customer, inactive_customer, other_channel_customer))
            db.flush()

            db.add_all(
                (
                    Product(sku="A-002", name="Producto inactivo", b2c_price=Decimal("900")),
                    Product(
                        sku="NO-PRICE",
                        name="Producto sin precio",
                        b2c_price=None,
                        default_route_id=self.base_packaging_route_id,
                        is_manufactured=True,
                    ),
                    Product(
                        sku="GLOBAL-INACTIVE",
                        name="Producto globalmente inactivo",
                        b2c_price=Decimal("950"),
                        active=False,
                    ),
                    Product(sku="OTHER", name="Producto de otro canal", b2c_price=Decimal("1000")),
                    Product(
                        sku="INACTIVE-CUSTOMER",
                        name="Producto de cliente inactivo",
                        b2c_price=Decimal("1100"),
                    ),
                    B2BCustomerProduct(
                        customer_id=duplicate_customer.id,
                        sku="A-001",
                        description="Té limón duplicado en otro cliente",
                        distributor_price=Decimal("2000"),
                        active=True,
                    ),
                    B2BCustomerProduct(
                        customer_id=duplicate_customer.id,
                        sku="NO-PRICE",
                        description="Producto sin precio",
                        distributor_price=Decimal("800"),
                        active=True,
                    ),
                    B2BCustomerProduct(
                        customer_id=duplicate_customer.id,
                        sku="GLOBAL-INACTIVE",
                        description="Producto globalmente inactivo",
                        distributor_price=Decimal("750"),
                        active=True,
                    ),
                    B2BCustomerProduct(
                        customer_id=other_channel_customer.id,
                        sku="OTHER",
                        description="Producto de otro canal",
                        distributor_price=Decimal("700"),
                        active=True,
                    ),
                    B2BCustomerProduct(
                        customer_id=inactive_customer.id,
                        sku="INACTIVE-CUSTOMER",
                        description="Producto de cliente inactivo",
                        distributor_price=Decimal("600"),
                        active=True,
                    ),
                )
            )

        with main.SessionLocal() as db:
            product_prices_before = dict(db.query(Product.sku, Product.b2c_price).all())
            product_states_before = dict(db.query(Product.sku, Product.active).all())
            catalog_states_before = dict(
                db.query(B2BCustomerProduct.id, B2BCustomerProduct.active).all()
            )

        response = self._export_updated_catalog()
        self.assertEqual(
            response.media_type,
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        self.assertEqual(
            response.headers["content-disposition"],
            'attachment; filename="Catalogo_B2B_actualizado.xlsx"',
        )
        workbook = load_workbook(io.BytesIO(response.body), data_only=False)
        worksheet = workbook["Catálogo B2B"]
        self.assertEqual(
            tuple(cell.value for cell in worksheet[1]),
            UPDATED_B2B_CATALOG_HEADERS,
        )
        self.assertEqual(worksheet.freeze_panes, "A2")
        self.assertEqual(worksheet.auto_filter.ref, f"A1:F{worksheet.max_row}")

        report_rows = list(worksheet.iter_rows(min_row=2, values_only=True))
        self.assertEqual([row[0] for row in report_rows].count("A-001"), 1)
        self.assertEqual([row[0] for row in report_rows], ["A-001", "NO-PRICE"])
        rows_by_sku = {row[0]: row for row in report_rows}
        self.assertEqual(set(rows_by_sku), {"A-001", "NO-PRICE"})
        self.assertEqual(rows_by_sku["A-001"][1], "Té limón")
        self.assertEqual(rows_by_sku["A-001"][2], "Infusiones")
        self.assertEqual(rows_by_sku["A-001"][3], 3200)
        self.assertAlmostEqual(rows_by_sku["A-001"][4], 0.268)
        self.assertAlmostEqual(rows_by_sku["A-001"][5], 2342.4)
        self.assertEqual(rows_by_sku["NO-PRICE"][3], MISSING_B2C_PRICE_LABEL)
        self.assertIsNone(rows_by_sku["NO-PRICE"][5])

        with main.SessionLocal() as db:
            self.assertEqual(
                dict(db.query(Product.sku, Product.b2c_price).all()),
                product_prices_before,
            )
            self.assertEqual(
                dict(db.query(Product.sku, Product.active).all()),
                product_states_before,
            )
            self.assertEqual(
                dict(db.query(B2BCustomerProduct.id, B2BCustomerProduct.active).all()),
                catalog_states_before,
            )

    def test_updated_catalog_filters_packaged_bulk_and_excludes_shipping_services(self) -> None:
        with main.SessionLocal.begin() as db:
            packaging_route = Route(
                code="R-PACKAGING-TEST",
                name="Envasado de prueba",
                process_type="packaging",
                active=True,
            )
            bulk_route = Route(
                code="R-BULK-TEST",
                name="Proceso a granel de prueba",
                process_type="dehydration",
                active=True,
            )
            db.add_all((packaging_route, bulk_route))
            db.flush()

            products = (
                Product(
                    sku="PACKAGED",
                    name="Producto envasado",
                    b2c_price=Decimal("1000"),
                    default_route_id=packaging_route.id,
                    is_manufactured=True,
                ),
                Product(
                    sku="BULK-ROUTE",
                    name="Producto a granel fabricado",
                    b2c_price=Decimal("1200"),
                    default_route_id=bulk_route.id,
                    is_manufactured=True,
                ),
                Product(
                    sku="BULK-PURCHASED",
                    name="Producto a granel comprado",
                    b2c_price=Decimal("1400"),
                    is_purchased_product=True,
                ),
                Product(
                    sku="MANUFACTURED-NO-ROUTE",
                    name="Producto fabricado sin ruta",
                    b2c_price=Decimal("1600"),
                    is_manufactured=True,
                ),
                *(
                    Product(
                        sku=sku,
                        name=f"Servicio de envío {index}",
                        b2c_price=Decimal("1800"),
                        default_route_id=(
                            packaging_route.id
                            if index == 1
                            else bulk_route.id if index == 2 else None
                        ),
                        is_manufactured=index == 4,
                        is_purchased_product=index == 3,
                    )
                    for index, sku in enumerate(sorted(EXCLUDED_SERVICE_SKUS), start=1)
                ),
            )
            db.add_all(products)
            db.flush()
            db.add_all(
                B2BCustomerProduct(
                    customer_id=self.customer_a_id,
                    sku=product.sku,
                    description=product.name,
                    distributor_price=Decimal("900"),
                    active=True,
                )
                for product in products
            )

        packaged_workbook = load_workbook(
            io.BytesIO(self._export_updated_catalog(product_scope=PRODUCT_SCOPE_PACKAGED).body),
            data_only=False,
        )
        packaged_skus = {
            row[0]
            for row in packaged_workbook["Catálogo B2B"].iter_rows(min_row=2, values_only=True)
        }
        self.assertEqual(packaged_skus, {"A-001", "PACKAGED"})

        bulk_workbook = load_workbook(
            io.BytesIO(self._export_updated_catalog(product_scope=PRODUCT_SCOPE_BULK).body),
            data_only=False,
        )
        bulk_skus = {
            row[0]
            for row in bulk_workbook["Catálogo B2B"].iter_rows(min_row=2, values_only=True)
        }
        self.assertEqual(
            bulk_skus,
            {"BULK-ROUTE", "BULK-PURCHASED", "MANUFACTURED-NO-ROUTE"},
        )

        all_workbook = load_workbook(
            io.BytesIO(self._export_updated_catalog(product_scope=PRODUCT_SCOPE_ALL).body),
            data_only=False,
        )
        all_skus = {
            row[0]
            for row in all_workbook["Catálogo B2B"].iter_rows(min_row=2, values_only=True)
        }
        self.assertEqual(
            all_skus,
            {
                "A-001",
                "PACKAGED",
                "BULK-ROUTE",
                "BULK-PURCHASED",
                "MANUFACTURED-NO-ROUTE",
            },
        )
        review_worksheet = all_workbook["Revisión"]
        self.assertEqual(tuple(cell.value for cell in review_worksheet[1]), REVIEW_HEADERS)
        review_skus = {
            row[0] for row in review_worksheet.iter_rows(min_row=2, values_only=True)
        }
        self.assertEqual(review_skus, set())
        for excluded_sku in EXCLUDED_SERVICE_SKUS:
            self.assertNotIn(excluded_sku, packaged_skus)
            self.assertNotIn(excluded_sku, bulk_skus)
            self.assertNotIn(excluded_sku, all_skus)
            self.assertNotIn(excluded_sku, review_skus)

    def test_non_packaging_route_is_classified_as_bulk(self) -> None:
        with main.SessionLocal.begin() as db:
            sterilized_jars_route = Route(
                code="R ESTERILIZADO FRASCOS",
                name="Esterilizado Frascos",
                process_type="dehydration",
                active=True,
            )
            db.add(sterilized_jars_route)
            db.flush()
            product = Product(
                sku="STERILIZED-JARS",
                name="Producto con ruta de esterilizado",
                b2c_price=Decimal("2000"),
                default_route_id=sterilized_jars_route.id,
                is_manufactured=True,
            )
            db.add(product)
            db.flush()
            db.add(
                B2BCustomerProduct(
                    customer_id=self.customer_a_id,
                    sku=product.sku,
                    description=product.name,
                    distributor_price=Decimal("1500"),
                    active=True,
                )
            )

        workbook = load_workbook(
            io.BytesIO(self._export_updated_catalog(product_scope=PRODUCT_SCOPE_BULK).body),
            data_only=False,
        )
        catalog_skus = {
            row[0]
            for row in workbook["Catálogo B2B"].iter_rows(min_row=2, values_only=True)
        }
        review_skus = {
            row[0] for row in workbook["Revisión"].iter_rows(min_row=2, values_only=True)
        }
        self.assertIn("STERILIZED-JARS", catalog_skus)
        self.assertNotIn("STERILIZED-JARS", review_skus)

    def test_updated_catalog_rejects_missing_or_ambiguous_canonical_skus(self) -> None:
        with main.SessionLocal.begin() as db:
            db.add(
                B2BCustomerProduct(
                    customer_id=self.customer_a_id,
                    sku="MISSING-PRODUCT",
                    description="Sin producto canónico",
                    distributor_price=Decimal("1000"),
                    active=True,
                )
            )
        with self.assertRaises(HTTPException) as missing_error:
            self._export_updated_catalog()
        self.assertEqual(missing_error.exception.status_code, 400)
        self.assertIn("MISSING-PRODUCT", str(missing_error.exception.detail))

        with main.SessionLocal.begin() as db:
            db.query(B2BCustomerProduct).filter(
                B2BCustomerProduct.sku == "MISSING-PRODUCT"
            ).delete()
            db.add_all(
                (
                    Product(sku="Case-SKU", name="Producto A", b2c_price=Decimal("100")),
                    Product(sku="case-sku", name="Producto B", b2c_price=Decimal("100")),
                    B2BCustomerProduct(
                        customer_id=self.customer_a_id,
                        sku="Case-SKU",
                        description="Producto A",
                        distributor_price=Decimal("90"),
                        active=True,
                    ),
                    B2BCustomerProduct(
                        customer_id=self.customer_a_id,
                        sku="case-sku",
                        description="Producto B",
                        distributor_price=Decimal("90"),
                        active=True,
                    ),
                )
            )
        with self.assertRaises(HTTPException) as ambiguous_error:
            self._export_updated_catalog()
        self.assertEqual(ambiguous_error.exception.status_code, 400)
        self.assertIn("ambiguos", str(ambiguous_error.exception.detail))

    def test_updated_catalog_validates_discount_and_requires_sales_view(self) -> None:
        with self.assertRaises(HTTPException) as invalid_discount:
            self._export_updated_catalog("100.01")
        self.assertEqual(invalid_discount.exception.status_code, 400)

        with self.assertRaises(HTTPException) as unauthorized:
            self._export_updated_catalog("10", {"product.view"})
        self.assertEqual(unauthorized.exception.status_code, 403)

        with self.assertRaises(HTTPException) as invalid_scope:
            self._export_updated_catalog(product_scope="not-a-scope")
        self.assertEqual(invalid_scope.exception.status_code, 400)

        self.assertEqual(
            sum(
                getattr(route, "path", "") == "/b2b/customers/catalog-updated.xlsx"
                for route in main.app.routes
            ),
            1,
        )

    def test_pdf_endpoint_exists_and_admin_and_sales_manager_can_download(self) -> None:
        self.assertEqual(
            sum(
                getattr(route, "path", "") == "/b2b/customers/{customer_id}/products/export.pdf"
                for route in main.app.routes
            ),
            1,
        )
        for role_code in ("admin", "sales_manager"):
            with self.subTest(role_code=role_code):
                response = self._export_customer_pdf(self.customer_a_id, self._role_permissions(role_code))
                self.assertEqual(response.media_type, "application/pdf")
                self.assertTrue(response.body.startswith(b"%PDF"))
                self.assertGreater(len(response.body), 1500)
                self.assertEqual(
                    response.headers["content-disposition"],
                    'attachment; filename="Catalogo_Cafe_Cliente_A.pdf"',
                )

    def test_pdf_exports_only_active_products_without_modifying_catalog(self) -> None:
        captured_products: list[B2BCustomerProduct] = []

        def build_pdf(_customer, products, *, logo_bytes=None):
            captured_products.extend(products)
            return b"%PDF-active-products-test"

        with mock.patch.object(main, "build_b2b_customer_catalog_pdf", side_effect=build_pdf):
            response = self._export_customer_pdf(
                self.customer_a_id,
                self._role_permissions("sales_manager"),
            )

        self.assertTrue(response.body.startswith(b"%PDF"))
        self.assertEqual([product.sku for product in captured_products], ["A-001"])
        self.assertTrue(all(product.active for product in captured_products))
        with main.SessionLocal() as db:
            states = dict(
                db.query(B2BCustomerProduct.sku, B2BCustomerProduct.active)
                .filter(B2BCustomerProduct.customer_id == self.customer_a_id)
                .all()
            )
        self.assertEqual(states, {"A-001": True, "A-002": False})

    def test_pdf_with_only_inactive_products_uses_empty_catalog(self) -> None:
        with main.SessionLocal.begin() as db:
            db.query(B2BCustomerProduct).filter(
                B2BCustomerProduct.customer_id == self.customer_a_id
            ).update({B2BCustomerProduct.active: False}, synchronize_session=False)

        with mock.patch.object(
            main,
            "build_b2b_customer_catalog_pdf",
            return_value=b"%PDF-empty-active-catalog-test",
        ) as build_pdf:
            response = self._export_customer_pdf(
                self.customer_a_id,
                self._role_permissions("admin"),
            )

        self.assertTrue(response.body.startswith(b"%PDF"))
        self.assertEqual(list(build_pdf.call_args.args[1]), [])

    def test_pdf_filename_is_readable_and_sanitized(self) -> None:
        self.assertEqual(
            safe_b2b_catalog_pdf_filename('Ángel / Azul: "Centro"'),
            "Catalogo_Angel_Azul_Centro.pdf",
        )
        self.assertEqual(safe_b2b_catalog_pdf_filename("  "), "Catalogo_Cliente.pdf")

    def test_pdf_endpoint_rejects_user_without_sales_view(self) -> None:
        with self.assertRaises(HTTPException) as raised:
            self._export_customer_pdf(self.customer_a_id, {"product.view"})
        self.assertEqual(raised.exception.status_code, 403)

    def test_pdf_handles_customer_without_channel_and_empty_catalog(self) -> None:
        with main.SessionLocal.begin() as db:
            empty_customer = B2BCustomer(
                customer_name="Cliente sin canal ni catálogo",
                address="Dirección extensa, con comas, tildes y suficiente texto para probar ajuste de línea",
            )
            db.add(empty_customer)
            db.flush()
            empty_customer_id = empty_customer.id

        response = self._export_customer_pdf(empty_customer_id, self._role_permissions("admin"))
        self.assertTrue(response.body.startswith(b"%PDF"))
        self.assertGreater(len(response.body), 1500)

    def test_pdf_currency_helper_uses_colones_and_rounds_to_whole_colones(self) -> None:
        self.assertEqual(format_crc_catalog_price(Decimal("3200.0000")), "₡3,200")
        self.assertEqual(format_crc_catalog_price(Decimal("2350.4900")), "₡2,350")
        self.assertEqual(format_crc_catalog_price(Decimal("20900.5000")), "₡20,901")

    def test_pdf_attempts_database_logo_and_survives_unreadable_logo(self) -> None:
        with main.SessionLocal.begin() as db:
            db.add(
                AppSettings(
                    id=main.BRANDING_SETTINGS_ID,
                    company_name="Green Corner",
                    logo_bytes=b"not-a-decodable-image",
                    logo_mime_type="image/png",
                )
            )

        with mock.patch("app.services.b2b_catalog_export_service._build_logo_flowable", return_value=None) as logo:
            response = self._export_customer_pdf(self.customer_a_id, self._role_permissions("admin"))
        logo.assert_called_once_with(b"not-a-decodable-image")
        self.assertTrue(response.body.startswith(b"%PDF"))

    def test_add_form_supports_manual_and_b2c_discount_pricing(self) -> None:
        catalog = Path("app/templates/b2b_customer_products.html").read_text(encoding="utf-8")
        self.assertIn("Método de precio", catalog)
        self.assertIn('value="manual"', catalog)
        self.assertIn("Digitar precio manual", catalog)
        self.assertIn('value="b2c_discount"', catalog)
        self.assertIn("Calcular por descuento sobre B2C", catalog)
        self.assertIn('id="b2c_discount_percent"', catalog)
        self.assertIn("Descuento sobre precio B2C (%)", catalog)
        self.assertIn("Precio sugerido al por mayor", catalog)
        self.assertIn("suggestedPrice = b2cPrice * (1 - (discountPercent / 100))", catalog)
        self.assertIn("distributorPriceInput.value = suggestedPrice.toFixed(2)", catalog)
        self.assertIn("El precio final de este campo es el único valor que se guarda.", catalog)

    def test_discount_mode_saves_final_distributor_price_not_discount(self) -> None:
        request = self._request_with_permissions(self._role_permissions("sales_manager"))
        with main.SessionLocal() as db, mock.patch.object(main, "safe_log_audit_event"):
            response = main.add_b2b_customer_product(
                customer_id=self.customer_a_id,
                request=request,
                sku="A-DISCOUNT",
                description="Precio final redondeado",
                distributor_price="2350",
                pricing_method="b2c_discount",
                discount_percent="26.8",
                active=True,
                db=db,
            )
            saved = (
                db.query(B2BCustomerProduct)
                .filter(
                    B2BCustomerProduct.customer_id == self.customer_a_id,
                    B2BCustomerProduct.sku == "A-DISCOUNT",
                )
                .one()
            )
        self.assertEqual(response.status_code, 303)
        self.assertEqual(saved.distributor_price, Decimal("2350.0000"))

    def test_admin_can_still_add_a_manual_price(self) -> None:
        request = self._request_with_permissions(self._role_permissions("admin"))
        with main.SessionLocal() as db, mock.patch.object(main, "safe_log_audit_event"):
            response = main.add_b2b_customer_product(
                customer_id=self.customer_a_id,
                request=request,
                sku="A-MANUAL",
                description="Precio manual",
                distributor_price="1800.50",
                pricing_method="manual",
                discount_percent="",
                active=True,
                db=db,
            )
            saved = db.query(B2BCustomerProduct).filter(B2BCustomerProduct.sku == "A-MANUAL").one()
        self.assertEqual(response.status_code, 303)
        self.assertEqual(saved.distributor_price, Decimal("1800.5000"))

    def test_user_without_catalog_edit_permission_cannot_add(self) -> None:
        request = self._request_with_permissions({"sales.view", "product.view"})
        with main.SessionLocal() as db:
            with self.assertRaises(HTTPException) as raised:
                main.add_b2b_customer_product(
                    customer_id=self.customer_a_id,
                    request=request,
                    sku="A-FORBIDDEN",
                    description="No debe guardarse",
                    distributor_price="1000",
                    pricing_method="manual",
                    discount_percent="",
                    active=True,
                    db=db,
                )
            self.assertEqual(raised.exception.status_code, 403)
            self.assertIsNone(
                db.query(B2BCustomerProduct).filter(B2BCustomerProduct.sku == "A-FORBIDDEN").one_or_none()
            )

    def test_add_pricing_validation_rejects_invalid_discount_and_nonpositive_final_price(self) -> None:
        with self.assertRaisesRegex(main.B2BValidationError, "between 0 and 100"):
            main._validate_b2b_catalog_add_pricing("b2c_discount", "100.01", "1000")
        with self.assertRaisesRegex(main.B2BValidationError, "greater than 0"):
            main._validate_b2b_catalog_add_pricing("manual", "", "0")


if __name__ == "__main__":
    unittest.main()
