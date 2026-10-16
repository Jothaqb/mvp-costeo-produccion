import io
import unittest
import zipfile
from datetime import date, datetime, timezone
from decimal import Decimal
from pathlib import Path
from unittest import mock

from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from app.database import Base
from app.models import (
    B2BCustomer,
    B2BCustomerProduct,
    B2BSalesOrder,
    B2BSalesOrderLine,
    Channel,
    Product,
)
from app.services.auth_service import BASE_PERMISSION_DEFINITIONS
from app.services import b2b_commercial_inactivity_review_service as review_service
from app.services.b2b_commercial_inactivity_review_service import (
    B2BCommercialInactivityReviewImportError,
    B2BCommercialInactivityReviewSecurityError,
    MANIFEST_SHEET_NAME,
    MAX_UPLOAD_BYTES,
    REVIEW_HEADERS,
    REVIEW_PERMISSION_DOWNLOAD,
    REVIEW_PERMISSION_EXECUTE,
    REVIEW_PERMISSION_VALIDATE,
    REVIEW_SHEET_NAME,
    build_b2b_inactivity_review_xlsx,
    build_inactivity_review_csrf_token,
    preview_b2b_inactivity_review_xlsx,
    validate_inactivity_review_csrf_token,
)
from app.services.b2b_commercial_inactivity_service import AUTHORIZED_B2B_CHANNEL_NAMES


AS_OF_DATE = date(2026, 10, 16)
SIGNING_SECRET = "review-test-secret-32-characters-minimum"
OLD_TIMESTAMP = datetime(2024, 1, 1, 8, 0, 0)


class B2BCommercialInactivityReviewTests(unittest.TestCase):
    def setUp(self):
        self.engine = create_engine(
            "sqlite://",
            connect_args={"check_same_thread": False},
            poolclass=StaticPool,
        )
        Base.metadata.create_all(self.engine)
        self.Session = sessionmaker(bind=self.engine, autocommit=False, autoflush=False)
        self.db = self.Session()
        self.channels = {}
        for name in AUTHORIZED_B2B_CHANNEL_NAMES:
            channel = Channel(
                name=name,
                active=True,
                applies_to_b2b=True,
                applies_to_b2c=False,
            )
            self.db.add(channel)
            self.db.flush()
            self.channels[name] = channel
        self.outside_channel = Channel(
            name="Canal fuera de alcance",
            active=True,
            applies_to_b2b=True,
            applies_to_b2c=False,
        )
        self.db.add(self.outside_channel)
        self.db.commit()

    def tearDown(self):
        self.db.close()
        Base.metadata.drop_all(self.engine)
        self.engine.dispose()

    def _customer(self, name="Cliente", *, channel=None):
        customer = B2BCustomer(
            customer_name=name,
            credit_days=0,
            active=True,
            channel_id=(channel or self.channels[AUTHORIZED_B2B_CHANNEL_NAMES[0]]).id,
            created_at=OLD_TIMESTAMP,
            updated_at=OLD_TIMESTAMP,
        )
        self.db.add(customer)
        self.db.flush()
        return customer

    def _catalog(self, customer, sku, *, description=None):
        self.db.add(
            Product(
                sku=sku,
                name=f"Producto {sku}",
                active=True,
                available_for_sale_gc=True,
                created_at=OLD_TIMESTAMP,
                updated_at=OLD_TIMESTAMP,
            )
        )
        row = B2BCustomerProduct(
            customer_id=customer.id,
            sku=sku,
            description=description or f"Producto {sku}",
            distributor_price=Decimal("1000"),
            active=True,
            created_at=OLD_TIMESTAMP,
            updated_at=OLD_TIMESTAMP,
        )
        self.db.add(row)
        self.db.flush()
        return row

    def _order(
        self,
        customer,
        order_number,
        sku,
        *,
        status="invoiced",
        invoice_date=None,
        line_total="1000",
    ):
        order = B2BSalesOrder(
            order_number=order_number,
            customer_id=customer.id,
            customer_name_snapshot=customer.customer_name,
            delivery_date=invoice_date or AS_OF_DATE,
            status=status,
            total_amount=Decimal(line_total),
            invoice_date=invoice_date if status == "invoiced" else None,
            created_at=OLD_TIMESTAMP,
            updated_at=OLD_TIMESTAMP,
        )
        order.lines = [
            B2BSalesOrderLine(
                line_number=1,
                sku_snapshot=sku,
                description_snapshot=f"Producto {sku}",
                unit_price_snapshot=Decimal(line_total),
                quantity=Decimal("1"),
                line_total=Decimal(line_total),
                created_at=OLD_TIMESTAMP,
                updated_at=OLD_TIMESTAMP,
            )
        ]
        self.db.add(order)
        self.db.flush()
        return order

    def _workbook(self, *, review_id="review-fixed-id"):
        self.db.commit()
        return build_b2b_inactivity_review_xlsx(
            self.db,
            as_of_date=AS_OF_DATE,
            review_id=review_id,
            generated_at=datetime(2026, 10, 16, 15, 0, tzinfo=timezone.utc),
            signing_secret=SIGNING_SECRET,
        )

    @staticmethod
    def _save_workbook(workbook):
        buffer = io.BytesIO()
        workbook.save(buffer)
        workbook.close()
        return buffer.getvalue()

    def test_download_contains_only_in_scope_inactivity_proposals_and_protected_rows(self):
        customer = self._customer('=HYPERLINK("https://invalid", "Cliente")')
        no_sale = self._catalog(customer, "NO-SALE", description="=1+1")
        old_sale = self._catalog(customer, "OLD-SALE")
        self._order(
            customer,
            "OLD-1",
            "OLD-SALE",
            invoice_date=date(2025, 1, 10),
            line_total="750.50",
        )
        recent = self._catalog(customer, "RECENT")
        self._order(
            customer,
            "RECENT-1",
            "RECENT",
            invoice_date=date(2026, 9, 1),
        )
        protected = self._catalog(customer, "OPEN")
        self._order(customer, "OPEN-1", "OPEN", status="draft", invoice_date=AS_OF_DATE)
        outside_customer = self._customer("Fuera", channel=self.outside_channel)
        outside = self._catalog(outside_customer, "OUTSIDE")

        review = self._workbook()
        workbook = load_workbook(io.BytesIO(review.content), data_only=False)
        worksheet = workbook[REVIEW_SHEET_NAME]
        rows_by_sku = {
            worksheet.cell(row_number, 3).value: row_number
            for row_number in range(2, worksheet.max_row + 1)
        }

        self.assertEqual(tuple(cell.value for cell in worksheet[1]), REVIEW_HEADERS)
        self.assertEqual(set(rows_by_sku), {"NO-SALE", "OLD-SALE", "OPEN"})
        self.assertNotIn(recent.id, [worksheet.cell(row, 9).value for row in rows_by_sku.values()])
        self.assertNotIn(outside.id, [worksheet.cell(row, 9).value for row in rows_by_sku.values()])
        self.assertEqual(
            worksheet.cell(rows_by_sku["NO-SALE"], 5).value,
            "Sin compra registrada en ERP",
        )
        self.assertEqual(worksheet.cell(rows_by_sku["OLD-SALE"], 6).value, "750.5000")
        self.assertIn("PROTEGIDO", worksheet.cell(rows_by_sku["OPEN"], 8).value)
        self.assertEqual(worksheet.freeze_panes, "A2")
        self.assertEqual(worksheet.auto_filter.ref, f"A1:J{worksheet.max_row}")
        self.assertEqual(worksheet.cell(rows_by_sku["NO-SALE"], 2).data_type, "s")
        self.assertEqual(worksheet.cell(rows_by_sku["NO-SALE"], 4).data_type, "s")
        self.assertEqual(workbook[MANIFEST_SHEET_NAME].sheet_state, "veryHidden")
        self.assertEqual(review.reviewable_count, 2)
        self.assertEqual(review.protected_count, 1)
        self.assertEqual(no_sale.active, True)
        self.assertEqual(protected.active, True)
        workbook.close()

    def test_import_kept_row_is_approved_and_deleted_row_is_maintained(self):
        customer = self._customer()
        first = self._catalog(customer, "FIRST")
        second = self._catalog(customer, "SECOND")
        review = self._workbook()
        workbook = load_workbook(io.BytesIO(review.content))
        worksheet = workbook[REVIEW_SHEET_NAME]
        second_row = next(
            row for row in range(2, worksheet.max_row + 1)
            if worksheet.cell(row, 3).value == "SECOND"
        )
        worksheet.delete_rows(second_row)
        uploaded = self._save_workbook(workbook)

        preview = preview_b2b_inactivity_review_xlsx(
            self.db,
            uploaded,
            current_as_of_date=AS_OF_DATE,
            signing_secret=SIGNING_SECRET,
        )

        self.assertEqual([item.catalog_product_id for item in preview.approved], [first.id])
        self.assertEqual([item.catalog_product_id for item in preview.maintained], [second.id])
        self.assertEqual(preview.protected, ())
        self.assertFalse(preview.plan_changed)
        self.assertTrue(first.active)
        self.assertTrue(second.active)

    def test_empty_visible_file_means_zero_approvals(self):
        customer = self._customer()
        first = self._catalog(customer, "FIRST")
        second = self._catalog(customer, "SECOND")
        review = self._workbook()
        workbook = load_workbook(io.BytesIO(review.content))
        worksheet = workbook[REVIEW_SHEET_NAME]
        worksheet.delete_rows(2, worksheet.max_row - 1)
        uploaded = self._save_workbook(workbook)

        preview = preview_b2b_inactivity_review_xlsx(
            self.db,
            uploaded,
            current_as_of_date=AS_OF_DATE,
            signing_secret=SIGNING_SECRET,
        )

        self.assertEqual(preview.approved, ())
        self.assertEqual(
            {item.catalog_product_id for item in preview.maintained},
            {first.id, second.id},
        )

    def test_open_order_row_remains_protected_when_kept_in_uploaded_file(self):
        customer = self._customer()
        protected = self._catalog(customer, "OPEN")
        self._order(customer, "OPEN-1", "OPEN", status="in_process")
        review = self._workbook()

        preview = preview_b2b_inactivity_review_xlsx(
            self.db,
            review.content,
            current_as_of_date=AS_OF_DATE,
            signing_secret=SIGNING_SECRET,
        )

        self.assertEqual(preview.approved, ())
        self.assertEqual([item.catalog_product_id for item in preview.protected], [protected.id])
        self.assertIn("pedido B2B abierto", preview.protected[0].detail)
        self.assertTrue(protected.active)

    def test_duplicate_new_and_edited_rows_are_rejected(self):
        customer = self._customer()
        self._catalog(customer, "FIRST")
        review = self._workbook()

        duplicate_workbook = load_workbook(io.BytesIO(review.content))
        duplicate_sheet = duplicate_workbook[REVIEW_SHEET_NAME]
        duplicate_sheet.append([cell.value for cell in duplicate_sheet[2]])
        duplicate_content = self._save_workbook(duplicate_workbook)
        with self.assertRaises(B2BCommercialInactivityReviewImportError) as duplicate_error:
            preview_b2b_inactivity_review_xlsx(
                self.db,
                duplicate_content,
                current_as_of_date=AS_OF_DATE,
                signing_secret=SIGNING_SECRET,
            )
        self.assertIn("duplicada", str(duplicate_error.exception))

        new_workbook = load_workbook(io.BytesIO(review.content))
        new_sheet = new_workbook[REVIEW_SHEET_NAME]
        copied = [cell.value for cell in new_sheet[2]]
        copied[8] = "999999"
        new_sheet.append(copied)
        new_content = self._save_workbook(new_workbook)
        with self.assertRaises(B2BCommercialInactivityReviewImportError) as new_error:
            preview_b2b_inactivity_review_xlsx(
                self.db,
                new_content,
                current_as_of_date=AS_OF_DATE,
                signing_secret=SIGNING_SECRET,
            )
        self.assertIn("no pertenecía", str(new_error.exception))

        edited_workbook = load_workbook(io.BytesIO(review.content))
        edited_sheet = edited_workbook[REVIEW_SHEET_NAME]
        edited_sheet.cell(2, 4).value = "Descripción alterada"
        edited_content = self._save_workbook(edited_workbook)
        with self.assertRaises(B2BCommercialInactivityReviewImportError) as edited_error:
            preview_b2b_inactivity_review_xlsx(
                self.db,
                edited_content,
                current_as_of_date=AS_OF_DATE,
                signing_secret=SIGNING_SECRET,
            )
        self.assertIn("fue alterada", str(edited_error.exception))

    def test_tampered_manifest_and_wrong_secret_are_rejected(self):
        customer = self._customer()
        self._catalog(customer, "FIRST")
        review = self._workbook()

        workbook = load_workbook(io.BytesIO(review.content))
        manifest = workbook[MANIFEST_SHEET_NAME]
        manifest.cell(4, 2).value = str(manifest.cell(4, 2).value) + "tampered"
        tampered_content = self._save_workbook(workbook)
        with self.assertRaises(B2BCommercialInactivityReviewImportError):
            preview_b2b_inactivity_review_xlsx(
                self.db,
                tampered_content,
                current_as_of_date=AS_OF_DATE,
                signing_secret=SIGNING_SECRET,
            )
        with self.assertRaises(B2BCommercialInactivityReviewImportError):
            preview_b2b_inactivity_review_xlsx(
                self.db,
                review.content,
                current_as_of_date=AS_OF_DATE,
                signing_secret="a-different-secret-that-is-long-enough",
            )

    def test_malformed_and_oversized_xlsx_files_are_rejected(self):
        for content in (b"not-an-xlsx", b"x" * (MAX_UPLOAD_BYTES + 1)):
            with self.assertRaises(B2BCommercialInactivityReviewImportError):
                preview_b2b_inactivity_review_xlsx(
                    self.db,
                    content,
                    current_as_of_date=AS_OF_DATE,
                    signing_secret=SIGNING_SECRET,
                )

        expanded = io.BytesIO()
        with zipfile.ZipFile(expanded, "w", compression=zipfile.ZIP_DEFLATED) as archive:
            archive.writestr("xl/worksheets/sheet1.xml", b"x" * 1_024)
        with mock.patch.object(review_service, "MAX_XLSX_UNCOMPRESSED_BYTES", 100):
            with self.assertRaises(B2BCommercialInactivityReviewImportError) as raised:
                preview_b2b_inactivity_review_xlsx(
                    self.db,
                    expanded.getvalue(),
                    current_as_of_date=AS_OF_DATE,
                    signing_secret=SIGNING_SECRET,
                )
        self.assertIn("tamaño descomprimido", str(raised.exception))

    def test_new_sale_after_download_protects_retained_assignment(self):
        customer = self._customer()
        catalog = self._catalog(customer, "FIRST")
        review = self._workbook()
        self._order(
            customer,
            "NEW-SALE",
            "FIRST",
            invoice_date=date(2026, 10, 10),
        )
        self.db.commit()

        preview = preview_b2b_inactivity_review_xlsx(
            self.db,
            review.content,
            current_as_of_date=AS_OF_DATE,
            signing_secret=SIGNING_SECRET,
        )

        self.assertEqual(preview.approved, ())
        self.assertEqual([item.catalog_product_id for item in preview.protected], [catalog.id])
        self.assertTrue(preview.plan_changed)
        self.assertIn("compra facturada", preview.protected[0].detail)
        self.assertTrue(catalog.active)

    def test_csrf_is_bound_to_session_and_uses_constant_time_validation(self):
        token = build_inactivity_review_csrf_token("session-one")
        validate_inactivity_review_csrf_token("session-one", token)
        with self.assertRaises(B2BCommercialInactivityReviewSecurityError):
            validate_inactivity_review_csrf_token("session-two", token)
        with self.assertRaises(B2BCommercialInactivityReviewSecurityError):
            validate_inactivity_review_csrf_token("session-one", "")

    def test_permissions_routes_and_disabled_execute_are_explicit(self):
        permission_codes = {row[0] for row in BASE_PERMISSION_DEFINITIONS}
        self.assertTrue(
            {
                REVIEW_PERMISSION_DOWNLOAD,
                REVIEW_PERMISSION_VALIDATE,
                REVIEW_PERMISSION_EXECUTE,
            }.issubset(permission_codes)
        )
        main_source = Path("app/main.py").read_text(encoding="utf-8")
        template = Path("app/templates/b2b_commercial_inactivity_review.html").read_text(
            encoding="utf-8"
        )
        self.assertIn(
            'require_permission(request, REVIEW_PERMISSION_DOWNLOAD)',
            main_source,
        )
        self.assertIn(
            'require_permission(request, REVIEW_PERMISSION_VALIDATE)',
            main_source,
        )
        self.assertIn("validate_inactivity_review_csrf_token(", main_source)
        self.assertIn("validate_role_permissions_request_origin(", main_source)
        self.assertIn(
            "await file.read(B2B_INACTIVITY_REVIEW_MAX_UPLOAD_BYTES + 1)",
            main_source,
        )
        self.assertNotIn('/b2b/commercial-inactivity-review/execute', main_source)
        self.assertIn("Confirmación no disponible", template)

    def test_secret_is_required_and_no_data_is_modified(self):
        customer = self._customer()
        catalog = self._catalog(customer, "FIRST")
        self.db.commit()
        with self.assertRaises(B2BCommercialInactivityReviewSecurityError):
            build_b2b_inactivity_review_xlsx(
                self.db,
                as_of_date=AS_OF_DATE,
                signing_secret="short",
            )
        self.db.refresh(customer)
        self.db.refresh(catalog)
        self.assertTrue(customer.active)
        self.assertTrue(catalog.active)


if __name__ == "__main__":
    unittest.main()
