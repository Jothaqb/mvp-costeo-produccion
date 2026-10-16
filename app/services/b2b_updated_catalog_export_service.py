from __future__ import annotations

from dataclasses import dataclass
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from app.models import B2BCustomer, B2BCustomerProduct, Channel, Product, ProductCategory, Route
from app.services.b2b_commercial_inactivity_service import AUTHORIZED_B2B_CHANNEL_NAMES


UPDATED_B2B_CATALOG_HEADERS = (
    "SKU",
    "Descripción del producto",
    "Categoría",
    "Precio B2C",
    "% Descuento aplicado",
    "Precio distribuidor calculado",
)

MONEY_QUANT = Decimal("0.0001")
PERCENT_BASE = Decimal("100")
MISSING_B2C_PRICE_LABEL = "Sin precio B2C"
_COLUMN_WIDTHS = (20, 46, 28, 20, 22, 30)
REVIEW_HEADERS = ("SKU", "Descripción del producto", "Motivo de revisión")
PRODUCT_SCOPE_PACKAGED = "packaged"
PRODUCT_SCOPE_BULK = "bulk"
PRODUCT_SCOPE_ALL = "all"
VALID_PRODUCT_SCOPES = frozenset(
    {
        PRODUCT_SCOPE_PACKAGED,
        PRODUCT_SCOPE_BULK,
        PRODUCT_SCOPE_ALL,
    }
)
_REVIEW_COLUMN_WIDTHS = (20, 46, 72)


class B2BUpdatedCatalogExportError(ValueError):
    pass


@dataclass(frozen=True)
class UpdatedB2BCatalogProduct:
    product_id: int
    sku: str
    product_name: str
    category_name: str | None
    b2c_price: Decimal | None
    active: bool
    default_route_id: int | None
    route_process_type: str | None


def parse_discount_percent(value: str | Decimal | int | float) -> Decimal:
    normalized = str(value).strip().replace(" ", "").replace(",", ".")
    if not normalized:
        raise B2BUpdatedCatalogExportError("El porcentaje de descuento es obligatorio.")
    try:
        discount = Decimal(normalized)
    except (InvalidOperation, ValueError) as exc:
        raise B2BUpdatedCatalogExportError(
            "El porcentaje de descuento debe ser un número entre 0 y 100."
        ) from exc
    if not discount.is_finite() or discount < 0 or discount > PERCENT_BASE:
        raise B2BUpdatedCatalogExportError(
            "El porcentaje de descuento debe estar entre 0 y 100."
        )
    return discount


def parse_product_scope(value: str) -> str:
    normalized = str(value or "").strip().lower()
    if normalized not in VALID_PRODUCT_SCOPES:
        raise B2BUpdatedCatalogExportError(
            "La presentación debe ser Envasados, A granel o Todos los productos comerciales."
        )
    return normalized


def build_updated_b2b_catalog_xlsx(
    db: Session,
    *,
    discount_percent: str | Decimal | int | float,
    product_scope: str = PRODUCT_SCOPE_ALL,
) -> bytes:
    discount = parse_discount_percent(discount_percent)
    normalized_scope = parse_product_scope(product_scope)
    products = _load_unique_catalog_products(db)
    classified_products: list[tuple[UpdatedB2BCatalogProduct, str]] = []
    review_products: list[UpdatedB2BCatalogProduct] = []
    for product in products:
        presentation = _commercial_presentation(product)
        if presentation is None:
            review_products.append(product)
        else:
            classified_products.append((product, presentation))
    selected_products = tuple(
        product
        for product, presentation in classified_products
        if normalized_scope == PRODUCT_SCOPE_ALL or presentation == normalized_scope
    )

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Catálogo B2B"
    worksheet.append(UPDATED_B2B_CATALOG_HEADERS)

    for product in selected_products:
        distributor_price = _discounted_price(product.b2c_price, discount)
        worksheet.append(
            (
                product.sku,
                product.product_name,
                product.category_name,
                product.b2c_price if product.b2c_price is not None else MISSING_B2C_PRICE_LABEL,
                discount / PERCENT_BASE,
                distributor_price,
            )
        )

    _format_worksheet(worksheet)
    review_worksheet = workbook.create_sheet("Revisión")
    review_worksheet.append(REVIEW_HEADERS)
    for product in review_products:
        review_worksheet.append(
            (
                product.sku,
                product.product_name,
                _review_reason(product),
            )
        )
    _format_review_worksheet(review_worksheet)
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _load_unique_catalog_products(db: Session) -> tuple[UpdatedB2BCatalogProduct, ...]:
    scoped_skus = {
        sku
        for (sku,) in (
            db.query(B2BCustomerProduct.sku)
            .join(B2BCustomer, B2BCustomer.id == B2BCustomerProduct.customer_id)
            .join(Channel, Channel.id == B2BCustomer.channel_id)
            .filter(
                B2BCustomer.active.is_(True),
                B2BCustomerProduct.active.is_(True),
                Channel.applies_to_b2b.is_(True),
                Channel.name.in_(AUTHORIZED_B2B_CHANNEL_NAMES),
            )
            .distinct()
            .all()
        )
    }

    rows = (
        db.query(
            Product.id,
            Product.sku,
            Product.name,
            ProductCategory.name,
            Product.b2c_price,
            Product.active,
            Product.default_route_id,
            Route.process_type,
        )
        .select_from(B2BCustomerProduct)
        .join(B2BCustomer, B2BCustomer.id == B2BCustomerProduct.customer_id)
        .join(Channel, Channel.id == B2BCustomer.channel_id)
        .join(Product, Product.sku == B2BCustomerProduct.sku)
        .outerjoin(ProductCategory, ProductCategory.id == Product.category_id)
        .outerjoin(Route, Route.id == Product.default_route_id)
        .filter(
            B2BCustomer.active.is_(True),
            B2BCustomerProduct.active.is_(True),
            Channel.applies_to_b2b.is_(True),
            Channel.name.in_(AUTHORIZED_B2B_CHANNEL_NAMES),
        )
        .distinct()
        .all()
    )

    mapped_skus = {row[1] for row in rows}
    missing_skus = sorted(scoped_skus - mapped_skus, key=str.casefold)
    if missing_skus:
        raise B2BUpdatedCatalogExportError(
            "No se pudo generar el catálogo porque hay SKU sin producto canónico: "
            + ", ".join(missing_skus)
            + "."
        )

    products = tuple(
        sorted(
            (
                UpdatedB2BCatalogProduct(
                    product_id=row[0],
                    sku=row[1],
                    product_name=row[2],
                    category_name=row[3],
                    b2c_price=row[4],
                    active=row[5],
                    default_route_id=row[6],
                    route_process_type=row[7],
                )
                for row in rows
            ),
            key=lambda product: (
                product.category_name is None,
                (product.category_name or "").casefold(),
                product.product_name.casefold(),
                product.sku.casefold(),
                product.product_id,
            ),
        )
    )
    _reject_ambiguous_skus(products)
    return tuple(product for product in products if product.active)


def _commercial_presentation(product: UpdatedB2BCatalogProduct) -> str | None:
    if product.route_process_type == "packaging":
        return PRODUCT_SCOPE_PACKAGED
    return None


def _review_reason(product: UpdatedB2BCatalogProduct) -> str:
    if product.default_route_id is None:
        return "Sin ruta predeterminada; la presentación comercial no puede determinarse con seguridad."
    return (
        "La ruta predeterminada representa un proceso distinto de packaging; "
        "la presentación comercial no puede determinarse con seguridad."
    )


def _reject_ambiguous_skus(products: tuple[UpdatedB2BCatalogProduct, ...]) -> None:
    product_ids_by_normalized_sku: dict[str, set[int]] = {}
    display_skus_by_normalized_sku: dict[str, set[str]] = {}
    for product in products:
        normalized_sku = product.sku.strip().casefold()
        product_ids_by_normalized_sku.setdefault(normalized_sku, set()).add(product.product_id)
        display_skus_by_normalized_sku.setdefault(normalized_sku, set()).add(product.sku)

    ambiguous = sorted(
        {
            " / ".join(sorted(display_skus_by_normalized_sku[normalized_sku], key=str.casefold))
            for normalized_sku, product_ids in product_ids_by_normalized_sku.items()
            if len(product_ids) > 1
        },
        key=str.casefold,
    )
    if ambiguous:
        raise B2BUpdatedCatalogExportError(
            "No se pudo generar el catálogo porque hay SKU ambiguos: "
            + ", ".join(ambiguous)
            + "."
        )


def _discounted_price(price: Decimal | None, discount: Decimal) -> Decimal | None:
    if price is None:
        return None
    return (price * (PERCENT_BASE - discount) / PERCENT_BASE).quantize(
        MONEY_QUANT,
        rounding=ROUND_HALF_UP,
    )


def _format_worksheet(worksheet: Any) -> None:
    header_fill = PatternFill(fill_type="solid", fgColor="1F4E3D")
    header_font = Font(color="FFFFFF", bold=True)
    last_column = get_column_letter(len(UPDATED_B2B_CATALOG_HEADERS))

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = f"A1:{last_column}{worksheet.max_row}"
    worksheet.row_dimensions[1].height = 32
    worksheet.sheet_view.showGridLines = False

    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for column_index, width in enumerate(_COLUMN_WIDTHS, start=1):
        worksheet.column_dimensions[get_column_letter(column_index)].width = width

    for row_index in range(2, worksheet.max_row + 1):
        worksheet.row_dimensions[row_index].height = 24
        for column_index in range(1, len(UPDATED_B2B_CATALOG_HEADERS) + 1):
            cell = worksheet.cell(row=row_index, column=column_index)
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            if isinstance(cell.value, str):
                # Prevent user-controlled product fields from being executed as formulas in Excel.
                cell.data_type = "s"

        worksheet.cell(row=row_index, column=1).number_format = "@"
        worksheet.cell(row=row_index, column=4).number_format = '₡ #,##0.00'
        worksheet.cell(row=row_index, column=5).number_format = "0.00%"
        worksheet.cell(row=row_index, column=6).number_format = '₡ #,##0.00'


def _format_review_worksheet(worksheet: Any) -> None:
    header_fill = PatternFill(fill_type="solid", fgColor="8A5A00")
    header_font = Font(color="FFFFFF", bold=True)
    last_column = get_column_letter(len(REVIEW_HEADERS))

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = f"A1:{last_column}{worksheet.max_row}"
    worksheet.row_dimensions[1].height = 32
    worksheet.sheet_view.showGridLines = False

    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for column_index, width in enumerate(_REVIEW_COLUMN_WIDTHS, start=1):
        worksheet.column_dimensions[get_column_letter(column_index)].width = width

    for row_index in range(2, worksheet.max_row + 1):
        worksheet.row_dimensions[row_index].height = 24
        for column_index in range(1, len(REVIEW_HEADERS) + 1):
            cell = worksheet.cell(row=row_index, column=column_index)
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            if isinstance(cell.value, str):
                cell.data_type = "s"
