from __future__ import annotations

import base64
import binascii
import hashlib
import hmac
import json
import os
import uuid
import zipfile
from dataclasses import dataclass
from datetime import date, datetime, timezone
from decimal import Decimal, InvalidOperation
from io import BytesIO
from typing import Any, Callable

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session, selectinload

from app.models import AuditLog, B2BCustomer, B2BCustomerProduct, B2BSalesOrder, User
from app.services.accounts_receivable_service import resolve_b2b_invoice_date
from app.services.audit_service import serialize_audit_payload
from app.services.b2b_commercial_inactivity_service import (
    B2BCommercialInactivityPlan,
    CatalogProductInactivityEvidence,
    InactivityConfig,
    build_b2b_commercial_inactivity_plan,
    current_costa_rica_date,
)


REVIEW_PERMISSION_DOWNLOAD = "b2b_commercial_inactivity.download"
REVIEW_PERMISSION_VALIDATE = "b2b_commercial_inactivity.validate"
REVIEW_PERMISSION_EXECUTE = "b2b_commercial_inactivity.execute"
REVIEW_SECRET_ENV = "B2B_INACTIVITY_REVIEW_SECRET"
REVIEW_WORKBOOK_VERSION = "b2b-inactivity-review-v1"
REVIEW_SHEET_NAME = "Propuestas"
MANIFEST_SHEET_NAME = "_manifiesto_firmado"
MAX_UPLOAD_BYTES = 10 * 1024 * 1024
MAX_XLSX_UNCOMPRESSED_BYTES = 50 * 1024 * 1024
MAX_XLSX_ZIP_ENTRIES = 1_000
MANIFEST_CHUNK_SIZE = 30_000
CSRF_PURPOSE = b"green-corner:b2b-commercial-inactivity-review:import:v1"
EXECUTION_CSRF_PURPOSE = b"green-corner:b2b-commercial-inactivity-review:execute:v1"
EXECUTION_TOKEN_PURPOSE = b"green-corner:b2b-commercial-inactivity-review:authorization:v1:"
EXECUTION_TOKEN_VERSION = "b2b-inactivity-review-execute-v1"
EXECUTION_TOKEN_MAX_AGE_SECONDS = 60 * 60
EXECUTION_ACCEPTANCE = "MANUAL_REVIEW_UNCERTIFIED_HISTORY_ACCEPTED"
HISTORY_COVERAGE_REASON = (
    "No existe una confirmación de cobertura histórica completa desde el inicio de la ventana."
)
REVIEW_HEADERS = (
    "Canal B2B",
    "Cliente",
    "SKU",
    "Descripción del producto",
    "Fecha de última compra facturada",
    "Monto neto facturado del SKU en la última compra",
    "Estado actual del producto en el catálogo",
    "Motivo de propuesta de inactivación",
    "Identificador interno de asignación producto-cliente",
    "Identificador único del proceso de revisión",
)


class B2BCommercialInactivityReviewError(RuntimeError):
    pass


class B2BCommercialInactivityReviewSecurityError(B2BCommercialInactivityReviewError):
    pass


class B2BCommercialInactivityReviewImportError(B2BCommercialInactivityReviewError):
    def __init__(self, errors: list[str] | tuple[str, ...]):
        self.errors = tuple(errors)
        super().__init__(" ".join(self.errors))


class B2BCommercialInactivityReviewExecutionError(B2BCommercialInactivityReviewError):
    pass


@dataclass(frozen=True)
class InactivityReviewRow:
    catalog_product_id: int
    customer_id: int
    channel_id: int
    channel_name: str
    customer_name: str
    sku: str
    description: str
    last_purchase_date: date | None
    last_purchase_amount: Decimal | None
    last_purchase_amount_note: str | None
    active_before: bool
    approval_allowed: bool
    protection_reasons: tuple[str, ...]
    proposal_reason: str
    visible_values: tuple[str, ...]


@dataclass(frozen=True)
class InactivityReviewWorkbook:
    content: bytes
    filename: str
    review_id: str
    plan_fingerprint: str
    as_of_date: date
    cutoff_date: date
    reviewable_count: int
    protected_count: int


@dataclass(frozen=True)
class InactivityReviewPreviewItem:
    catalog_product_id: int
    customer_id: int
    channel_id: int
    channel_name: str
    customer_name: str
    sku: str
    description: str
    decision: str
    detail: str


@dataclass(frozen=True)
class InactivityReviewGroupTotal:
    label: str
    approved: int
    maintained: int
    protected: int


@dataclass(frozen=True)
class InactivityReviewPreview:
    review_id: str
    generated_at_utc: str
    original_as_of_date: date
    current_as_of_date: date
    original_plan_fingerprint: str
    current_plan_fingerprint: str
    manifest_sha256: str
    plan_changed: bool
    approved: tuple[InactivityReviewPreviewItem, ...]
    maintained: tuple[InactivityReviewPreviewItem, ...]
    protected: tuple[InactivityReviewPreviewItem, ...]
    changes: tuple[str, ...]
    channel_totals: tuple[InactivityReviewGroupTotal, ...]
    customer_totals: tuple[InactivityReviewGroupTotal, ...]
    original_row_count: int
    uploaded_row_count: int

    @property
    def affected_customer_count(self) -> int:
        return len({item.customer_id for item in self.approved})


@dataclass(frozen=True)
class InactivityReviewExecutionResult:
    review_id: str
    plan_fingerprint: str
    deactivated: tuple[InactivityReviewPreviewItem, ...]

    @property
    def affected_customer_count(self) -> int:
        return len({item.customer_id for item in self.deactivated})


@dataclass(frozen=True)
class InactivityReviewApprovedClaim:
    catalog_product_id: int
    customer_id: int
    channel_id: int
    sku: str


@dataclass(frozen=True)
class InactivityReviewExecutionAuthorization:
    review_id: str
    manifest_sha256: str
    plan_fingerprint: str
    as_of_date: date
    issued_at_utc: datetime
    approved: tuple[InactivityReviewApprovedClaim, ...]


def build_inactivity_review_csrf_token(session_token: str | None) -> str:
    normalized = (session_token or "").strip()
    if not normalized:
        raise B2BCommercialInactivityReviewSecurityError(
            "No hay una sesión válida para generar la protección CSRF."
        )
    return hmac.new(normalized.encode("utf-8"), CSRF_PURPOSE, hashlib.sha256).hexdigest()


def validate_inactivity_review_csrf_token(
    session_token: str | None,
    submitted_token: str | None,
) -> None:
    submitted = (submitted_token or "").strip()
    if not submitted:
        raise B2BCommercialInactivityReviewSecurityError(
            "Falta el token CSRF del formulario."
        )
    expected = build_inactivity_review_csrf_token(session_token)
    if not hmac.compare_digest(submitted, expected):
        raise B2BCommercialInactivityReviewSecurityError(
            "El token CSRF no es válido para esta sesión."
        )


def build_inactivity_review_execute_csrf_token(session_token: str | None) -> str:
    normalized = (session_token or "").strip()
    if not normalized:
        raise B2BCommercialInactivityReviewSecurityError(
            "No hay una sesión válida para generar la protección CSRF de confirmación."
        )
    return hmac.new(
        normalized.encode("utf-8"),
        EXECUTION_CSRF_PURPOSE,
        hashlib.sha256,
    ).hexdigest()


def validate_inactivity_review_execute_csrf_token(
    session_token: str | None,
    submitted_token: str | None,
) -> None:
    submitted = (submitted_token or "").strip()
    if not submitted:
        raise B2BCommercialInactivityReviewSecurityError(
            "Falta el token CSRF de confirmación."
        )
    expected = build_inactivity_review_execute_csrf_token(session_token)
    if not hmac.compare_digest(submitted, expected):
        raise B2BCommercialInactivityReviewSecurityError(
            "El token CSRF de confirmación no es válido para esta sesión."
        )


def build_b2b_inactivity_review_xlsx(
    db: Session,
    *,
    as_of_date: date | None = None,
    review_id: str | None = None,
    generated_at: datetime | None = None,
    signing_secret: str | None = None,
) -> InactivityReviewWorkbook:
    effective_date = as_of_date or current_costa_rica_date()
    effective_review_id = review_id or str(uuid.uuid4())
    effective_generated_at = generated_at or datetime.now(timezone.utc)
    plan = build_b2b_commercial_inactivity_plan(
        db,
        config=InactivityConfig(as_of_date=effective_date, history_complete_since=None),
    )
    if not plan.channel_scope_is_valid:
        raise B2BCommercialInactivityReviewError(
            "No se puede generar la revisión porque los cuatro canales autorizados "
            "no se resolvieron exactamente."
        )
    rows = _build_review_rows(db, plan=plan, review_id=effective_review_id)
    manifest = {
        "version": REVIEW_WORKBOOK_VERSION,
        "review_id": effective_review_id,
        "generated_at_utc": effective_generated_at.astimezone(timezone.utc).isoformat(),
        "as_of_date": effective_date.isoformat(),
        "cutoff_date": plan.config.cutoff_date.isoformat(),
        "plan_fingerprint": plan.plan_fingerprint,
        "channel_scope_revision": plan.channel_scope_revision,
        "eligibility_policy_revision": plan.eligibility_policy_revision,
        "rows": [_manifest_row(row) for row in rows],
    }
    canonical_manifest = _canonical_json(manifest)
    signature = _sign_manifest(canonical_manifest, signing_secret=signing_secret)
    content = _render_workbook(rows, canonical_manifest=canonical_manifest, signature=signature)
    filename = f"Revision_Inactividad_B2B_{effective_date.isoformat()}_{effective_review_id[:8]}.xlsx"
    return InactivityReviewWorkbook(
        content=content,
        filename=filename,
        review_id=effective_review_id,
        plan_fingerprint=plan.plan_fingerprint,
        as_of_date=effective_date,
        cutoff_date=plan.config.cutoff_date,
        reviewable_count=sum(row.approval_allowed for row in rows),
        protected_count=sum(not row.approval_allowed for row in rows),
    )


def preview_b2b_inactivity_review_xlsx(
    db: Session,
    content: bytes,
    *,
    current_as_of_date: date | None = None,
    signing_secret: str | None = None,
) -> InactivityReviewPreview:
    if not content:
        raise B2BCommercialInactivityReviewImportError(["El archivo XLSX está vacío."])
    if len(content) > MAX_UPLOAD_BYTES:
        raise B2BCommercialInactivityReviewImportError(
            ["El archivo XLSX supera el máximo permitido de 10 MB."]
        )
    _validate_xlsx_container(content)
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=False)
    except Exception as exc:
        raise B2BCommercialInactivityReviewImportError(
            ["El archivo no es un XLSX válido o está dañado."]
        ) from exc
    try:
        manifest, manifest_sha256 = _read_and_validate_manifest(
            workbook,
            signing_secret=signing_secret,
        )
        retained_ids = _validate_visible_rows(workbook, manifest)
    finally:
        workbook.close()

    effective_current_date = current_as_of_date or current_costa_rica_date()
    current_plan = build_b2b_commercial_inactivity_plan(
        db,
        config=InactivityConfig(
            as_of_date=effective_current_date,
            history_complete_since=None,
        ),
    )
    original_rows = {
        int(row["catalog_product_id"]): row
        for row in manifest["rows"]
    }
    current_evidence = {
        row.catalog_product_id: row for row in current_plan.catalog_products
    }
    current_candidate_ids = {
        row.catalog_product_id
        for row in current_plan.catalog_products
        if _is_inactivity_proposal(row, current_plan)
    }

    approved: list[InactivityReviewPreviewItem] = []
    maintained: list[InactivityReviewPreviewItem] = []
    protected: list[InactivityReviewPreviewItem] = []
    changes: list[str] = []
    plan_changed = manifest["plan_fingerprint"] != current_plan.plan_fingerprint
    if plan_changed:
        changes.append(
            "El fingerprint del universo cambió desde la descarga; cualquier confirmación futura "
            "deberá usar una revisión recién generada y revalidada."
        )
    if manifest["as_of_date"] != effective_current_date.isoformat():
        changes.append(
            "La fecha operativa actual difiere de la fecha usada al generar el archivo."
        )

    for catalog_product_id, original in sorted(original_rows.items()):
        item_base = {
            "catalog_product_id": catalog_product_id,
            "customer_id": int(original["customer_id"]),
            "channel_id": int(original["channel_id"]),
            "channel_name": str(original["channel_name"]),
            "customer_name": str(original["customer_name"]),
            "sku": str(original["sku"]),
            "description": str(original["description"]),
        }
        original_approval_allowed = bool(original["approval_allowed"])
        if not original_approval_allowed:
            detail = " | ".join(original.get("protection_reasons") or ())
            protected.append(
                InactivityReviewPreviewItem(
                    **item_base,
                    decision="Protegido",
                    detail=detail or "La asignación no era autorizable al generar el archivo.",
                )
            )
            continue
        if catalog_product_id not in retained_ids:
            maintained.append(
                InactivityReviewPreviewItem(
                    **item_base,
                    decision="Mantener activo",
                    detail="La fila fue eliminada del Excel validado por el usuario.",
                )
            )
            continue

        evidence = current_evidence.get(catalog_product_id)
        current_blockers = _current_blockers(original, evidence, current_plan)
        if current_blockers:
            protected.append(
                InactivityReviewPreviewItem(
                    **item_base,
                    decision="Protegido por cambio posterior",
                    detail=" | ".join(current_blockers),
                )
            )
            changes.append(
                f"Asignación {catalog_product_id} ({original['customer_name']} / {original['sku']}): "
                + " | ".join(current_blockers)
            )
            continue
        approved.append(
            InactivityReviewPreviewItem(
                **item_base,
                decision="Aprobado para confirmar",
                detail=(
                    "La fila se conservó y continúa sin compras facturadas en la ventana; "
                    "todavía no se realizó ninguna inactivación."
                ),
            )
        )

    new_candidate_count = len(current_candidate_ids - set(original_rows))
    if new_candidate_count:
        changes.append(
            f"Hay {new_candidate_count} propuesta(s) nueva(s) que no pertenecen a este proceso de revisión."
        )

    all_items = (*approved, *maintained, *protected)
    return InactivityReviewPreview(
        review_id=str(manifest["review_id"]),
        generated_at_utc=str(manifest["generated_at_utc"]),
        original_as_of_date=date.fromisoformat(str(manifest["as_of_date"])),
        current_as_of_date=effective_current_date,
        original_plan_fingerprint=str(manifest["plan_fingerprint"]),
        current_plan_fingerprint=current_plan.plan_fingerprint,
        manifest_sha256=manifest_sha256,
        plan_changed=plan_changed,
        approved=tuple(approved),
        maintained=tuple(maintained),
        protected=tuple(protected),
        changes=tuple(changes),
        channel_totals=_group_totals(all_items, key="channel_name"),
        customer_totals=_group_totals(all_items, key="customer_name"),
        original_row_count=len(original_rows),
        uploaded_row_count=len(retained_ids),
    )


def build_b2b_inactivity_execution_token(
    preview: InactivityReviewPreview,
    *,
    session_token: str | None,
    issued_at: datetime | None = None,
    signing_secret: str | None = None,
) -> str:
    if not preview.approved:
        raise B2BCommercialInactivityReviewExecutionError(
            "No hay asignaciones aprobadas para confirmar."
        )
    if preview.plan_changed:
        raise B2BCommercialInactivityReviewExecutionError(
            "El plan cambió desde la descarga; genere y valide una nueva revisión."
        )
    normalized_session = (session_token or "").strip()
    if not normalized_session:
        raise B2BCommercialInactivityReviewSecurityError(
            "No hay una sesión válida para autorizar la confirmación."
        )
    secret = _review_secret(signing_secret)
    effective_issued_at = issued_at or datetime.now(timezone.utc)
    if effective_issued_at.tzinfo is None:
        effective_issued_at = effective_issued_at.replace(tzinfo=timezone.utc)
    payload = {
        "version": EXECUTION_TOKEN_VERSION,
        "review_id": preview.review_id,
        "manifest_sha256": preview.manifest_sha256,
        "plan_fingerprint": preview.current_plan_fingerprint,
        "as_of_date": preview.current_as_of_date.isoformat(),
        "issued_at_utc": effective_issued_at.astimezone(timezone.utc).isoformat(),
        "session_binding": _execution_session_binding(secret, normalized_session),
        "approved": [
            {
                "catalog_product_id": item.catalog_product_id,
                "customer_id": item.customer_id,
                "channel_id": item.channel_id,
                "sku": item.sku,
            }
            for item in sorted(preview.approved, key=lambda row: row.catalog_product_id)
        ],
    }
    canonical = _canonical_json(payload).encode("utf-8")
    signature = hmac.new(secret, EXECUTION_TOKEN_PURPOSE + canonical, hashlib.sha256).hexdigest()
    encoded = base64.urlsafe_b64encode(canonical).decode("ascii").rstrip("=")
    return f"{encoded}.{signature}"


def validate_b2b_inactivity_execution_token(
    token: str | None,
    *,
    session_token: str | None,
    current_time: datetime | None = None,
    signing_secret: str | None = None,
) -> InactivityReviewExecutionAuthorization:
    normalized_token = (token or "").strip()
    normalized_session = (session_token or "").strip()
    if not normalized_token or len(normalized_token) > 1_000_000:
        raise B2BCommercialInactivityReviewSecurityError(
            "La autorización firmada de confirmación no es válida."
        )
    if not normalized_session:
        raise B2BCommercialInactivityReviewSecurityError(
            "No hay una sesión válida para confirmar la revisión."
        )
    try:
        encoded, signature = normalized_token.split(".", 1)
        padded = encoded + "=" * (-len(encoded) % 4)
        canonical = base64.urlsafe_b64decode(padded.encode("ascii"))
    except (binascii.Error, ValueError, UnicodeEncodeError) as exc:
        raise B2BCommercialInactivityReviewSecurityError(
            "La autorización firmada de confirmación no es válida."
        ) from exc
    secret = _review_secret(signing_secret)
    expected_signature = hmac.new(
        secret,
        EXECUTION_TOKEN_PURPOSE + canonical,
        hashlib.sha256,
    ).hexdigest()
    if not signature or not hmac.compare_digest(signature, expected_signature):
        raise B2BCommercialInactivityReviewSecurityError(
            "La autorización de confirmación fue alterada o no pertenece a este ERP."
        )
    try:
        payload = json.loads(canonical.decode("utf-8"))
        issued_at = datetime.fromisoformat(str(payload["issued_at_utc"]))
        as_of_date = date.fromisoformat(str(payload["as_of_date"]))
        approved_payload = payload["approved"]
    except (KeyError, TypeError, ValueError, UnicodeDecodeError, json.JSONDecodeError) as exc:
        raise B2BCommercialInactivityReviewSecurityError(
            "La autorización firmada de confirmación está incompleta."
        ) from exc
    if payload.get("version") != EXECUTION_TOKEN_VERSION or not isinstance(approved_payload, list):
        raise B2BCommercialInactivityReviewSecurityError(
            "La versión de la autorización de confirmación no es compatible."
        )
    if issued_at.tzinfo is None:
        raise B2BCommercialInactivityReviewSecurityError(
            "La autorización de confirmación no contiene una fecha UTC válida."
        )
    effective_now = current_time or datetime.now(timezone.utc)
    if effective_now.tzinfo is None:
        effective_now = effective_now.replace(tzinfo=timezone.utc)
    age_seconds = (
        effective_now.astimezone(timezone.utc) - issued_at.astimezone(timezone.utc)
    ).total_seconds()
    if age_seconds < -300 or age_seconds > EXECUTION_TOKEN_MAX_AGE_SECONDS:
        raise B2BCommercialInactivityReviewSecurityError(
            "La autorización de confirmación venció; cargue nuevamente el Excel."
        )
    expected_binding = _execution_session_binding(secret, normalized_session)
    if not hmac.compare_digest(str(payload.get("session_binding") or ""), expected_binding):
        raise B2BCommercialInactivityReviewSecurityError(
            "La autorización de confirmación no pertenece a esta sesión."
        )
    review_id = str(payload.get("review_id") or "").strip()
    manifest_sha256 = str(payload.get("manifest_sha256") or "").strip()
    plan_fingerprint = str(payload.get("plan_fingerprint") or "").strip()
    if not review_id or not _is_sha256(manifest_sha256) or not _is_sha256(plan_fingerprint):
        raise B2BCommercialInactivityReviewSecurityError(
            "La autorización firmada no identifica una revisión válida."
        )
    approved: list[InactivityReviewApprovedClaim] = []
    approved_ids: set[int] = set()
    try:
        for raw_claim in approved_payload:
            claim = InactivityReviewApprovedClaim(
                catalog_product_id=int(raw_claim["catalog_product_id"]),
                customer_id=int(raw_claim["customer_id"]),
                channel_id=int(raw_claim["channel_id"]),
                sku=str(raw_claim["sku"]).strip(),
            )
            if claim.catalog_product_id <= 0 or claim.customer_id <= 0 or claim.channel_id <= 0:
                raise ValueError("Invalid claim identifier")
            if not claim.sku or claim.catalog_product_id in approved_ids:
                raise ValueError("Invalid or duplicate approved assignment")
            approved_ids.add(claim.catalog_product_id)
            approved.append(claim)
    except (KeyError, TypeError, ValueError) as exc:
        raise B2BCommercialInactivityReviewSecurityError(
            "La autorización contiene asignaciones inválidas o duplicadas."
        ) from exc
    if not approved:
        raise B2BCommercialInactivityReviewExecutionError(
            "La revisión no contiene asignaciones aprobadas para confirmar."
        )
    return InactivityReviewExecutionAuthorization(
        review_id=review_id,
        manifest_sha256=manifest_sha256,
        plan_fingerprint=plan_fingerprint,
        as_of_date=as_of_date,
        issued_at_utc=issued_at.astimezone(timezone.utc),
        approved=tuple(approved),
    )


def execute_b2b_inactivity_review(
    db: Session,
    *,
    authorization_token: str | None,
    session_token: str | None,
    acceptance: str | None,
    acting_user_id: int,
    request_path: str,
    request_method: str,
    ip_address: str | None,
    current_as_of_date: date | None = None,
    current_time: datetime | None = None,
    signing_secret: str | None = None,
    before_commit: Callable[[Session], None] | None = None,
) -> InactivityReviewExecutionResult:
    if (acceptance or "").strip() != EXECUTION_ACCEPTANCE:
        raise B2BCommercialInactivityReviewExecutionError(
            "Debe aceptar expresamente que esta es una revisión comercial manual con histórico no certificado."
        )
    authorization = validate_b2b_inactivity_execution_token(
        authorization_token,
        session_token=session_token,
        current_time=current_time,
        signing_secret=signing_secret,
    )
    effective_date = current_as_of_date or current_costa_rica_date()
    if effective_date != authorization.as_of_date:
        raise B2BCommercialInactivityReviewExecutionError(
            "La fecha operativa cambió; genere y valide una nueva revisión."
        )
    try:
        actor = (
            db.query(User)
            .filter(User.id == acting_user_id, User.is_active.is_(True))
            .one_or_none()
        )
        if actor is None:
            raise B2BCommercialInactivityReviewExecutionError(
                "Se requiere un usuario activo para registrar la confirmación."
            )
        approved_ids = sorted(claim.catalog_product_id for claim in authorization.approved)
        customer_ids = sorted({claim.customer_id for claim in authorization.approved})
        locked_customers = (
            db.query(B2BCustomer)
            .filter(B2BCustomer.id.in_(customer_ids))
            .with_for_update()
            .all()
        )
        locked_assignments = (
            db.query(B2BCustomerProduct)
            .filter(B2BCustomerProduct.id.in_(approved_ids))
            .with_for_update()
            .all()
        )
        if len(locked_customers) != len(customer_ids) or len(locked_assignments) != len(approved_ids):
            raise B2BCommercialInactivityReviewExecutionError(
                "Una asignación o cliente aprobado ya no existe; genere una nueva revisión."
            )

        db.expire_all()
        plan = build_b2b_commercial_inactivity_plan(
            db,
            config=InactivityConfig(as_of_date=effective_date, history_complete_since=None),
        )
        if not plan.channel_scope_is_valid:
            raise B2BCommercialInactivityReviewExecutionError(
                "Los cuatro canales autorizados ya no se resuelven exactamente."
            )
        if not hmac.compare_digest(plan.plan_fingerprint, authorization.plan_fingerprint):
            raise B2BCommercialInactivityReviewExecutionError(
                "El fingerprint cambió desde la aprobación; genere y valide una nueva revisión."
            )
        current_evidence = {
            row.catalog_product_id: row for row in plan.catalog_products
        }
        validation_errors: list[str] = []
        for claim in authorization.approved:
            evidence = current_evidence.get(claim.catalog_product_id)
            original = {
                "customer_id": claim.customer_id,
                "channel_id": claim.channel_id,
                "sku": claim.sku,
            }
            blockers = _current_blockers(original, evidence, plan)
            if evidence is None or not _is_inactivity_proposal(evidence, plan):
                blockers = (*blockers, "La asignación ya no es una propuesta de inactividad.")
            if blockers:
                validation_errors.append(
                    f"Asignación {claim.catalog_product_id}: " + " | ".join(dict.fromkeys(blockers))
                )
        if validation_errors:
            raise B2BCommercialInactivityReviewExecutionError(
                "La revisión cambió y no puede ejecutarse parcialmente. "
                + " ".join(validation_errors)
            )

        assignments_by_id = {
            row.id: row
            for row in (
                db.query(B2BCustomerProduct)
                .filter(B2BCustomerProduct.id.in_(approved_ids))
                .order_by(B2BCustomerProduct.id)
                .all()
            )
        }
        deactivated: list[InactivityReviewPreviewItem] = []
        for claim in sorted(authorization.approved, key=lambda row: row.catalog_product_id):
            assignment = assignments_by_id[claim.catalog_product_id]
            evidence = current_evidence[claim.catalog_product_id]
            if not assignment.active:
                raise B2BCommercialInactivityReviewExecutionError(
                    f"La asignación {assignment.id} ya no está activa."
                )
            assignment.active = False
            audit_detail = {
                "active": False,
                "result": "deactivated",
                "review_id": authorization.review_id,
                "manifest_sha256": authorization.manifest_sha256,
                "plan_fingerprint": authorization.plan_fingerprint,
                "customer_id": evidence.customer_id,
                "customer_name": evidence.customer_name,
                "sku": evidence.sku,
                "manual_history_acknowledgement": True,
            }
            db.add(
                AuditLog(
                    user_id=actor.id,
                    username=actor.username,
                    module="sales",
                    action="b2b_customer_product_inactivity_review_deactivated",
                    entity_type="b2b_customer_product",
                    entity_id=str(assignment.id),
                    entity_label=f"{evidence.customer_name} - {evidence.sku}"[:255],
                    old_values=serialize_audit_payload({"active": True}),
                    new_values=serialize_audit_payload(audit_detail),
                    request_path=request_path,
                    method=request_method,
                    ip_address=((ip_address or "")[:255] or None),
                    notes=(
                        "Confirmación manual mediante revisión Excel firmada; "
                        "la cobertura histórica no se marcó como certificada."
                    ),
                )
            )
            deactivated.append(
                InactivityReviewPreviewItem(
                    catalog_product_id=assignment.id,
                    customer_id=evidence.customer_id,
                    channel_id=evidence.channel_id,
                    channel_name=evidence.channel_name,
                    customer_name=evidence.customer_name,
                    sku=evidence.sku,
                    description=evidence.description,
                    decision="Inactivado",
                    detail="Asignación producto–cliente desactivada mediante confirmación manual.",
                )
            )
        db.flush()
        remaining_active = (
            db.query(B2BCustomerProduct.id)
            .filter(
                B2BCustomerProduct.id.in_(approved_ids),
                B2BCustomerProduct.active.is_(True),
            )
            .count()
        )
        if remaining_active:
            raise B2BCommercialInactivityReviewExecutionError(
                "No fue posible verificar todas las inactivaciones aprobadas."
            )
        if before_commit is not None:
            before_commit(db)
        db.commit()
        return InactivityReviewExecutionResult(
            review_id=authorization.review_id,
            plan_fingerprint=authorization.plan_fingerprint,
            deactivated=tuple(deactivated),
        )
    except Exception:
        db.rollback()
        raise


def _execution_session_binding(secret: bytes, session_token: str) -> str:
    return hmac.new(
        secret,
        b"session:" + session_token.encode("utf-8"),
        hashlib.sha256,
    ).hexdigest()


def _is_sha256(value: str) -> bool:
    if len(value) != 64:
        return False
    try:
        int(value, 16)
    except ValueError:
        return False
    return True


def _validate_xlsx_container(content: bytes) -> None:
    try:
        with zipfile.ZipFile(BytesIO(content)) as archive:
            entries = archive.infolist()
    except (OSError, zipfile.BadZipFile) as exc:
        raise B2BCommercialInactivityReviewImportError(
            ["El archivo no es un XLSX válido o está dañado."]
        ) from exc
    if len(entries) > MAX_XLSX_ZIP_ENTRIES:
        raise B2BCommercialInactivityReviewImportError(
            ["El archivo XLSX contiene demasiados componentes internos."]
        )
    if sum(entry.file_size for entry in entries) > MAX_XLSX_UNCOMPRESSED_BYTES:
        raise B2BCommercialInactivityReviewImportError(
            ["El archivo XLSX supera el tamaño descomprimido permitido de 50 MB."]
        )


def _build_review_rows(
    db: Session,
    *,
    plan: B2BCommercialInactivityPlan,
    review_id: str,
) -> tuple[InactivityReviewRow, ...]:
    purchase_amounts, unavailable_amounts = _latest_purchase_amounts(db, plan=plan)
    rows: list[InactivityReviewRow] = []
    for evidence in plan.catalog_products:
        if not _is_inactivity_proposal(evidence, plan):
            continue
        protection_reasons = tuple(
            reason for reason in evidence.reasons if reason != HISTORY_COVERAGE_REASON
        )
        approval_allowed = not protection_reasons
        if approval_allowed:
            proposal_reason = (
                "Aprobación manual requerida: sin compras facturadas del SKU durante los últimos "
                "12 meses; la cobertura histórica todavía no está certificada."
            )
        else:
            proposal_reason = "PROTEGIDO — " + " | ".join(protection_reasons)
        amount_key = (
            evidence.customer_id,
            evidence.sku,
            evidence.last_purchase_date,
        )
        last_purchase_amount = purchase_amounts.get(amount_key)
        amount_note: str | None = None
        if evidence.last_purchase_date is None:
            amount_note = "No disponible: sin compra registrada en ERP."
        elif amount_key in unavailable_amounts or last_purchase_amount is None:
            amount_note = (
                "No disponible: el importe de la última compra no pudo calcularse con seguridad."
            )
        last_purchase_display = (
            evidence.last_purchase_date.isoformat()
            if evidence.last_purchase_date is not None
            else "Sin compra registrada en ERP"
        )
        amount_display = (
            format(last_purchase_amount, "f")
            if last_purchase_amount is not None and amount_note is None
            else amount_note or "No disponible"
        )
        visible_values = (
            evidence.channel_name,
            evidence.customer_name,
            evidence.sku,
            evidence.description,
            last_purchase_display,
            amount_display,
            "Activo" if evidence.active_before else "Inactivo",
            proposal_reason,
            str(evidence.catalog_product_id),
            review_id,
        )
        rows.append(
            InactivityReviewRow(
                catalog_product_id=evidence.catalog_product_id,
                customer_id=evidence.customer_id,
                channel_id=evidence.channel_id,
                channel_name=evidence.channel_name,
                customer_name=evidence.customer_name,
                sku=evidence.sku,
                description=evidence.description,
                last_purchase_date=evidence.last_purchase_date,
                last_purchase_amount=last_purchase_amount,
                last_purchase_amount_note=amount_note,
                active_before=evidence.active_before,
                approval_allowed=approval_allowed,
                protection_reasons=protection_reasons,
                proposal_reason=proposal_reason,
                visible_values=visible_values,
            )
        )
    rows.sort(key=lambda row: (row.channel_name.casefold(), row.customer_name.casefold(), row.sku))
    return tuple(rows)


def _is_inactivity_proposal(
    evidence: CatalogProductInactivityEvidence,
    plan: B2BCommercialInactivityPlan,
) -> bool:
    return bool(
        evidence.active_before
        and (
            evidence.last_purchase_date is None
            or evidence.last_purchase_date < plan.config.cutoff_date
        )
    )


def _latest_purchase_amounts(
    db: Session,
    *,
    plan: B2BCommercialInactivityPlan,
) -> tuple[dict[tuple[int, str, date | None], Decimal], set[tuple[int, str, date | None]]]:
    customer_ids = sorted({row.customer_id for row in plan.catalog_products})
    if not customer_ids:
        return {}, set()
    orders = (
        db.query(B2BSalesOrder)
        .filter(
            B2BSalesOrder.status == "invoiced",
            B2BSalesOrder.customer_id.in_(customer_ids),
        )
        .options(selectinload(B2BSalesOrder.lines))
        .order_by(B2BSalesOrder.id)
        .all()
    )
    latest_dates = {
        (row.customer_id, row.sku): row.last_purchase_date
        for row in plan.catalog_products
        if row.last_purchase_date is not None
    }
    amounts: dict[tuple[int, str, date | None], Decimal] = {}
    unavailable: set[tuple[int, str, date | None]] = set()
    for order in orders:
        invoice_date, _source = resolve_b2b_invoice_date(order)
        if invoice_date > plan.config.as_of_date:
            continue
        for line in order.lines:
            sku = (line.sku_snapshot or "").strip()
            if not sku or (line.quantity or Decimal("0")) <= Decimal("0"):
                continue
            if latest_dates.get((order.customer_id, sku)) != invoice_date:
                continue
            key = (order.customer_id, sku, invoice_date)
            try:
                line_amount = Decimal(str(line.line_total))
            except (InvalidOperation, TypeError, ValueError):
                unavailable.add(key)
                continue
            amounts[key] = amounts.get(key, Decimal("0")) + line_amount
    for key in unavailable:
        amounts.pop(key, None)
    return amounts, unavailable


def _manifest_row(row: InactivityReviewRow) -> dict[str, Any]:
    return {
        "catalog_product_id": row.catalog_product_id,
        "customer_id": row.customer_id,
        "channel_id": row.channel_id,
        "channel_name": row.channel_name,
        "customer_name": row.customer_name,
        "sku": row.sku,
        "description": row.description,
        "last_purchase_date": (
            row.last_purchase_date.isoformat() if row.last_purchase_date is not None else None
        ),
        "last_purchase_amount": (
            format(row.last_purchase_amount, "f")
            if row.last_purchase_amount is not None
            else None
        ),
        "last_purchase_amount_note": row.last_purchase_amount_note,
        "active_before": row.active_before,
        "approval_allowed": row.approval_allowed,
        "protection_reasons": list(row.protection_reasons),
        "proposal_reason": row.proposal_reason,
        "visible_values": list(row.visible_values),
    }


def _render_workbook(
    rows: tuple[InactivityReviewRow, ...],
    *,
    canonical_manifest: str,
    signature: str,
) -> bytes:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = REVIEW_SHEET_NAME
    worksheet.append(REVIEW_HEADERS)
    for row in rows:
        worksheet.append(row.visible_values)
        for cell in worksheet[worksheet.max_row]:
            cell.data_type = "s"
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = f"A1:J{max(1, worksheet.max_row)}"
    header_fill = PatternFill("solid", fgColor="1F4E78")
    protected_fill = PatternFill("solid", fgColor="FCE4D6")
    for cell in worksheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row_number, row in enumerate(rows, start=2):
        if not row.approval_allowed:
            for cell in worksheet[row_number]:
                cell.fill = protected_fill
        for cell in worksheet[row_number]:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    widths = (30, 32, 16, 38, 24, 28, 20, 72, 22, 38)
    for index, width in enumerate(widths, start=1):
        worksheet.column_dimensions[get_column_letter(index)].width = width
    worksheet.sheet_view.showGridLines = False

    manifest_sheet = workbook.create_sheet(MANIFEST_SHEET_NAME)
    manifest_sheet.append(("version", REVIEW_WORKBOOK_VERSION))
    manifest_sheet.append(("signature_sha256", signature))
    chunks = [
        canonical_manifest[index : index + MANIFEST_CHUNK_SIZE]
        for index in range(0, len(canonical_manifest), MANIFEST_CHUNK_SIZE)
    ] or [""]
    manifest_sheet.append(("chunk_count", str(len(chunks))))
    for index, chunk in enumerate(chunks, start=1):
        manifest_sheet.append((f"chunk_{index:04d}", chunk))
    manifest_sheet.sheet_state = "veryHidden"
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _read_and_validate_manifest(
    workbook,
    *,
    signing_secret: str | None,
) -> tuple[dict[str, Any], str]:
    if REVIEW_SHEET_NAME not in workbook.sheetnames or MANIFEST_SHEET_NAME not in workbook.sheetnames:
        raise B2BCommercialInactivityReviewImportError(
            ["El archivo no contiene las hojas originales de la revisión comercial."]
        )
    manifest_sheet = workbook[MANIFEST_SHEET_NAME]
    values = {
        str(row[0].value): row[1].value
        for row in manifest_sheet.iter_rows(min_row=1, max_col=2)
        if row[0].value is not None
    }
    if str(values.get("version") or "") != REVIEW_WORKBOOK_VERSION:
        raise B2BCommercialInactivityReviewImportError(
            ["La versión del archivo de revisión no es compatible."]
        )
    try:
        chunk_count = int(str(values.get("chunk_count") or "0"))
    except ValueError as exc:
        raise B2BCommercialInactivityReviewImportError(
            ["El manifiesto firmado del archivo es inválido."]
        ) from exc
    canonical_manifest = "".join(
        str(values.get(f"chunk_{index:04d}") or "")
        for index in range(1, chunk_count + 1)
    )
    signature = str(values.get("signature_sha256") or "")
    expected_signature = _sign_manifest(canonical_manifest, signing_secret=signing_secret)
    if not signature or not hmac.compare_digest(signature, expected_signature):
        raise B2BCommercialInactivityReviewImportError(
            ["El manifiesto del archivo fue alterado o no pertenece a este ERP."]
        )
    try:
        manifest = json.loads(canonical_manifest)
    except json.JSONDecodeError as exc:
        raise B2BCommercialInactivityReviewImportError(
            ["El manifiesto firmado no contiene JSON válido."]
        ) from exc
    required = {
        "version",
        "review_id",
        "generated_at_utc",
        "as_of_date",
        "cutoff_date",
        "plan_fingerprint",
        "channel_scope_revision",
        "eligibility_policy_revision",
        "rows",
    }
    missing = sorted(required - set(manifest))
    if missing or not isinstance(manifest.get("rows"), list):
        raise B2BCommercialInactivityReviewImportError(
            ["El manifiesto firmado está incompleto."]
        )
    if manifest.get("version") != REVIEW_WORKBOOK_VERSION:
        raise B2BCommercialInactivityReviewImportError(
            ["El manifiesto pertenece a otra versión del proceso de revisión."]
        )
    return manifest, hashlib.sha256(canonical_manifest.encode("utf-8")).hexdigest()


def _validate_visible_rows(workbook, manifest: dict[str, Any]) -> set[int]:
    worksheet = workbook[REVIEW_SHEET_NAME]
    raw_rows = list(worksheet.iter_rows(values_only=True))
    if not raw_rows:
        raise B2BCommercialInactivityReviewImportError(
            ["La hoja de propuestas no contiene encabezados."]
        )
    detected_headers = tuple(_cell_text(value) for value in raw_rows[0][: len(REVIEW_HEADERS)])
    if detected_headers != REVIEW_HEADERS:
        raise B2BCommercialInactivityReviewImportError(
            [
                "Los encabezados de la hoja fueron modificados. Deben mantenerse exactamente como "
                "en el archivo descargado."
            ]
        )
    manifest_rows: dict[int, dict[str, Any]] = {}
    for row in manifest["rows"]:
        try:
            row_id = int(row["catalog_product_id"])
        except (KeyError, TypeError, ValueError) as exc:
            raise B2BCommercialInactivityReviewImportError(
                ["El manifiesto contiene un identificador de asignación inválido."]
            ) from exc
        if row_id in manifest_rows:
            raise B2BCommercialInactivityReviewImportError(
                ["El manifiesto contiene identificadores duplicados."]
            )
        manifest_rows[row_id] = row

    errors: list[str] = []
    retained_ids: set[int] = set()
    for row_number, raw_row in enumerate(raw_rows[1:], start=2):
        values = tuple(_cell_text(value) for value in raw_row[: len(REVIEW_HEADERS)])
        if not any(values):
            continue
        if len(raw_row) > len(REVIEW_HEADERS) and any(
            _cell_text(value) for value in raw_row[len(REVIEW_HEADERS) :]
        ):
            errors.append(f"Fila {row_number}: contiene columnas adicionales no autorizadas.")
            continue
        raw_id = values[8]
        try:
            row_id = int(raw_id)
        except ValueError:
            errors.append(f"Fila {row_number}: identificador interno inválido.")
            continue
        if row_id in retained_ids:
            errors.append(f"Fila {row_number}: la asignación {row_id} está duplicada.")
            continue
        original = manifest_rows.get(row_id)
        if original is None:
            errors.append(
                f"Fila {row_number}: la asignación {row_id} no pertenecía al proceso original."
            )
            continue
        if values[9] != str(manifest["review_id"]):
            errors.append(f"Fila {row_number}: pertenece a otro proceso de revisión.")
            continue
        expected_values = tuple(_cell_text(value) for value in original["visible_values"])
        if values != expected_values:
            errors.append(
                f"Fila {row_number}: fue alterada. Solo se permite conservar o eliminar filas completas."
            )
            continue
        retained_ids.add(row_id)
    if errors:
        raise B2BCommercialInactivityReviewImportError(errors)
    return retained_ids


def _current_blockers(
    original: dict[str, Any],
    evidence: CatalogProductInactivityEvidence | None,
    current_plan: B2BCommercialInactivityPlan,
) -> tuple[str, ...]:
    if evidence is None:
        return (
            "La asignación ya no existe, cambió a un canal fuera del alcance o dejó de formar "
            "parte del universo autorizado.",
        )
    blockers: list[str] = []
    if evidence.customer_id != int(original["customer_id"]):
        blockers.append("Cambió el cliente asociado.")
    if evidence.channel_id != int(original["channel_id"]):
        blockers.append("Cambió el canal B2B del cliente.")
    if evidence.sku != str(original["sku"]):
        blockers.append("Cambió el SKU de la asignación.")
    if not evidence.active_before:
        blockers.append("La asignación ya no está activa.")
    if (
        evidence.last_purchase_date is not None
        and evidence.last_purchase_date >= current_plan.config.cutoff_date
    ):
        blockers.append("Ahora existe una compra facturada dentro de los últimos 12 meses.")
    blockers.extend(
        reason for reason in evidence.reasons if reason != HISTORY_COVERAGE_REASON
    )
    return tuple(dict.fromkeys(blockers))


def _group_totals(
    items: tuple[InactivityReviewPreviewItem, ...],
    *,
    key: str,
) -> tuple[InactivityReviewGroupTotal, ...]:
    totals: dict[str, dict[str, int]] = {}
    for item in items:
        label = str(getattr(item, key))
        bucket = totals.setdefault(label, {"approved": 0, "maintained": 0, "protected": 0})
        if item.decision.startswith("Aprobado"):
            bucket["approved"] += 1
        elif item.decision.startswith("Mantener"):
            bucket["maintained"] += 1
        else:
            bucket["protected"] += 1
    return tuple(
        InactivityReviewGroupTotal(label=label, **counts)
        for label, counts in sorted(totals.items(), key=lambda entry: entry[0].casefold())
    )


def _canonical_json(payload: dict[str, Any]) -> str:
    return json.dumps(payload, ensure_ascii=False, sort_keys=True, separators=(",", ":"))


def _sign_manifest(canonical_manifest: str, *, signing_secret: str | None) -> str:
    secret = _review_secret(signing_secret)
    return hmac.new(secret, canonical_manifest.encode("utf-8"), hashlib.sha256).hexdigest()


def _review_secret(explicit_secret: str | None) -> bytes:
    secret = explicit_secret if explicit_secret is not None else os.getenv(REVIEW_SECRET_ENV, "")
    normalized = secret.strip()
    if len(normalized) < 32:
        raise B2BCommercialInactivityReviewSecurityError(
            f"Configure {REVIEW_SECRET_ENV} con un secreto estable de al menos 32 caracteres."
        )
    return normalized.encode("utf-8")


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, datetime):
        return value.isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return str(value).strip()
