from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime, time, timedelta, timezone
from decimal import Decimal
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import and_, case, func, not_, or_
from sqlalchemy.orm import Session, joinedload, selectinload

from app.models import B2BCustomer, B2BSalesOrder
from app.services.accounts_receivable_service import (
    HISTORICAL_CSV_IMPORT_PREFIX,
    STATUS_OVERDUE,
    build_accounts_receivable_dashboard,
)


B2B_CUSTOMER_EXPORT_HEADERS = (
    "Nombre del cliente",
    "Canal B2B",
    "Estado del cliente",
    "Razón social",
    "Cédula jurídica",
    "Dirección",
    "Provincia",
    "Cantón",
    "Distrito",
    "Teléfono",
    "Forma de pago",
    "Días de crédito",
    "Estado del último acuerdo",
    "¿Tenemos acuerdo firmado?",
)

NO_REGISTERED_SIGNATURE = "Sin firma registrada"
_TEXT_COLUMN_INDEXES = (5, 10)
_SALES_COLUMN_INDEXES = (15, 16)
_LAST_INVOICE_DATE_COLUMN_INDEX = 17
_OVERDUE_AMOUNT_COLUMN_INDEX = 19
_COSTA_RICA_TIMEZONE = timezone(timedelta(hours=-6), name="America/Costa_Rica")
_COLUMN_WIDTHS = (
    30,
    30,
    20,
    34,
    22,
    46,
    18,
    20,
    20,
    20,
    38,
    16,
    28,
    28,
    20,
    24,
    24,
    24,
    22,
    20,
)


@dataclass(frozen=True)
class B2BCustomerSalesSummary:
    previous_year_sales: Decimal
    current_year_sales: Decimal
    last_invoice_date: date | None


@dataclass(frozen=True)
class B2BCustomerOverdueSummary:
    overdue_invoice_count: int
    overdue_amount: Decimal
    max_days_overdue: int


def current_b2b_customer_export_datetime() -> datetime:
    return datetime.now(_COSTA_RICA_TIMEZONE)


def b2b_customer_export_report_date(generated_at: datetime) -> date:
    return _costa_rica_datetime(generated_at).date()


def b2b_customer_export_headers(report_year: int) -> tuple[str, ...]:
    return (
        *B2B_CUSTOMER_EXPORT_HEADERS,
        f"Ventas {report_year - 1}",
        f"Ventas acumuladas {report_year}",
        "Fecha última facturación",
        "Cantidad de facturas vencidas",
        "Monto total vencido",
        "Días de mayor atraso",
    )


def build_b2b_customers_export_xlsx(
    db: Session,
    *,
    generated_at: datetime | None = None,
) -> bytes:
    effective_generated_at = _costa_rica_datetime(generated_at)
    report_date = b2b_customer_export_report_date(effective_generated_at)
    headers = b2b_customer_export_headers(report_date.year)
    customers = (
        db.query(B2BCustomer)
        .options(
            joinedload(B2BCustomer.channel_record),
            selectinload(B2BCustomer.commercial_agreements),
        )
        .order_by(B2BCustomer.customer_name, B2BCustomer.id)
        .all()
    )
    sales_by_customer_id = _sales_summary_by_customer(db, report_date=report_date)
    overdue_by_customer_id = _overdue_summary_by_customer(db, report_date=report_date)

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Clientes B2B"
    worksheet.append(headers)

    for customer in customers:
        worksheet.append(
            _customer_export_row(
                customer,
                sales_by_customer_id.get(
                    customer.id,
                    B2BCustomerSalesSummary(
                        previous_year_sales=Decimal("0"),
                        current_year_sales=Decimal("0"),
                        last_invoice_date=None,
                    ),
                ),
                overdue_by_customer_id.get(
                    customer.id,
                    B2BCustomerOverdueSummary(
                        overdue_invoice_count=0,
                        overdue_amount=Decimal("0"),
                        max_days_overdue=0,
                    ),
                ),
            )
        )

    _format_worksheet(worksheet, headers=headers)
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _customer_export_row(
    customer: B2BCustomer,
    sales_summary: B2BCustomerSalesSummary,
    overdue_summary: B2BCustomerOverdueSummary,
) -> tuple[Any, ...]:
    agreements = list(customer.commercial_agreements)
    latest_agreement = max(
        agreements,
        key=lambda agreement: (agreement.version_number, agreement.id or 0),
        default=None,
    )
    registered_signed_agreements = [
        agreement
        for agreement in agreements
        if agreement.status in {"signed", "archived"} and agreement.signed_at is not None
    ]
    latest_registered_signature = max(
        registered_signed_agreements,
        key=lambda agreement: (
            agreement.signed_at,
            agreement.version_number,
            agreement.id or 0,
        ),
        default=None,
    )

    return (
        customer.customer_name,
        customer.channel_record.name if customer.channel_record is not None else None,
        "Activo" if customer.active else "Inactivo",
        customer.legal_name,
        _text_or_none(customer.legal_id),
        customer.address,
        customer.province,
        customer.canton,
        customer.district,
        _text_or_none(customer.phone),
        customer.default_payment_terms,
        customer.credit_days,
        latest_agreement.status if latest_agreement is not None else None,
        "Sí" if latest_registered_signature is not None else NO_REGISTERED_SIGNATURE,
        sales_summary.previous_year_sales,
        sales_summary.current_year_sales,
        sales_summary.last_invoice_date,
        overdue_summary.overdue_invoice_count,
        overdue_summary.overdue_amount,
        overdue_summary.max_days_overdue,
    )


def _costa_rica_datetime(value: datetime | None) -> datetime:
    if value is None:
        return datetime.now(_COSTA_RICA_TIMEZONE)
    if value.tzinfo is None:
        return value.replace(tzinfo=_COSTA_RICA_TIMEZONE)
    return value.astimezone(_COSTA_RICA_TIMEZONE)


def _sales_summary_by_customer(
    db: Session,
    *,
    report_date: date,
) -> dict[int, B2BCustomerSalesSummary]:
    previous_year_start = date(report_date.year - 1, 1, 1)
    current_year_start = date(report_date.year, 1, 1)
    report_end_exclusive = report_date + timedelta(days=1)

    invoice_date_source = B2BSalesOrder.invoice_date.is_not(None)
    invoiced_at_source = and_(
        B2BSalesOrder.invoice_date.is_(None),
        B2BSalesOrder.invoiced_at.is_not(None),
    )
    historical_marker = (
        func.substr(
            func.coalesce(B2BSalesOrder.observations, ""),
            1,
            len(HISTORICAL_CSV_IMPORT_PREFIX),
        )
        == HISTORICAL_CSV_IMPORT_PREFIX
    )
    historical_delivery_source = and_(
        B2BSalesOrder.invoice_date.is_(None),
        B2BSalesOrder.invoiced_at.is_(None),
        historical_marker,
    )
    created_at_source = and_(
        B2BSalesOrder.invoice_date.is_(None),
        B2BSalesOrder.invoiced_at.is_(None),
        not_(historical_marker),
    )

    def effective_date_in_range(start: date, end_exclusive: date):
        start_datetime = datetime.combine(start, time.min)
        end_datetime = datetime.combine(end_exclusive, time.min)
        return or_(
            and_(
                invoice_date_source,
                B2BSalesOrder.invoice_date >= start,
                B2BSalesOrder.invoice_date < end_exclusive,
            ),
            and_(
                invoiced_at_source,
                B2BSalesOrder.invoiced_at >= start_datetime,
                B2BSalesOrder.invoiced_at < end_datetime,
            ),
            and_(
                historical_delivery_source,
                B2BSalesOrder.delivery_date >= start,
                B2BSalesOrder.delivery_date < end_exclusive,
            ),
            and_(
                created_at_source,
                B2BSalesOrder.created_at >= start_datetime,
                B2BSalesOrder.created_at < end_datetime,
            ),
        )

    previous_year_condition = effective_date_in_range(previous_year_start, current_year_start)
    current_year_condition = effective_date_in_range(current_year_start, report_end_exclusive)
    zero = Decimal("0")

    rows = (
        db.query(
            B2BSalesOrder.customer_id.label("customer_id"),
            func.coalesce(
                func.sum(
                    case(
                        (previous_year_condition, B2BSalesOrder.total_amount),
                        else_=zero,
                    )
                ),
                zero,
            ).label("previous_year_sales"),
            func.coalesce(
                func.sum(
                    case(
                        (current_year_condition, B2BSalesOrder.total_amount),
                        else_=zero,
                    )
                ),
                zero,
            ).label("current_year_sales"),
            func.max(
                case(
                    (invoice_date_source, B2BSalesOrder.invoice_date),
                    else_=None,
                )
            ).label("latest_invoice_date"),
            func.max(
                case(
                    (invoiced_at_source, B2BSalesOrder.invoiced_at),
                    else_=None,
                )
            ).label("latest_invoiced_at"),
            func.max(
                case(
                    (historical_delivery_source, B2BSalesOrder.delivery_date),
                    else_=None,
                )
            ).label("latest_historical_delivery_date"),
            func.max(
                case(
                    (created_at_source, B2BSalesOrder.created_at),
                    else_=None,
                )
            ).label("latest_created_at"),
        )
        .filter(B2BSalesOrder.status == "invoiced")
        .group_by(B2BSalesOrder.customer_id)
        .all()
    )

    summaries: dict[int, B2BCustomerSalesSummary] = {}
    for row in rows:
        latest_dates = [
            value
            for value in (
                _date_value(row.latest_invoice_date),
                _date_value(row.latest_invoiced_at),
                _date_value(row.latest_historical_delivery_date),
                _date_value(row.latest_created_at),
            )
            if value is not None
        ]
        summaries[row.customer_id] = B2BCustomerSalesSummary(
            previous_year_sales=Decimal(str(row.previous_year_sales or 0)),
            current_year_sales=Decimal(str(row.current_year_sales or 0)),
            last_invoice_date=max(latest_dates, default=None),
        )
    return summaries


def _overdue_summary_by_customer(
    db: Session,
    *,
    report_date: date,
) -> dict[int, B2BCustomerOverdueSummary]:
    dashboard = build_accounts_receivable_dashboard(db, today=report_date)
    counts_by_customer_id: dict[int, int] = {}
    amounts_by_customer_id: dict[int, Decimal] = {}
    max_days_by_customer_id: dict[int, int] = {}

    for row in dashboard.rows:
        if row.status_code != STATUS_OVERDUE:
            continue
        days_overdue = -(row.days_remaining or 0)
        counts_by_customer_id[row.customer_id] = counts_by_customer_id.get(row.customer_id, 0) + 1
        amounts_by_customer_id[row.customer_id] = (
            amounts_by_customer_id.get(row.customer_id, Decimal("0")) + row.pending_amount
        )
        max_days_by_customer_id[row.customer_id] = max(
            max_days_by_customer_id.get(row.customer_id, 0),
            days_overdue,
        )

    return {
        customer_id: B2BCustomerOverdueSummary(
            overdue_invoice_count=count,
            overdue_amount=amounts_by_customer_id[customer_id],
            max_days_overdue=max_days_by_customer_id[customer_id],
        )
        for customer_id, count in counts_by_customer_id.items()
    }


def _text_or_none(value: Any) -> str | None:
    if value is None:
        return None
    return str(value)


def _date_value(value: date | datetime | None) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    return value


def _format_worksheet(worksheet, *, headers: tuple[str, ...]) -> None:
    header_fill = PatternFill(fill_type="solid", fgColor="1F4E3D")
    header_font = Font(color="FFFFFF", bold=True)

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{worksheet.max_row}"
    worksheet.row_dimensions[1].height = 34
    worksheet.sheet_view.showGridLines = False

    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for column_index, width in enumerate(_COLUMN_WIDTHS, start=1):
        worksheet.column_dimensions[get_column_letter(column_index)].width = width

    for row_index in range(2, worksheet.max_row + 1):
        worksheet.row_dimensions[row_index].height = 28
        for column_index in range(1, len(headers) + 1):
            cell = worksheet.cell(row=row_index, column=column_index)
            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True,
            )
            if isinstance(cell.value, str):
                # Prevent user-controlled values such as "=1+1" from becoming
                # executable formulas when the workbook is opened in Excel.
                cell.data_type = "s"
        for column_index in _TEXT_COLUMN_INDEXES:
            worksheet.cell(row=row_index, column=column_index).number_format = "@"
        for column_index in _SALES_COLUMN_INDEXES:
            worksheet.cell(row=row_index, column=column_index).number_format = "₡ #,##0.00"
        worksheet.cell(
            row=row_index,
            column=_LAST_INVOICE_DATE_COLUMN_INDEX,
        ).number_format = "dd/mm/yyyy"
        worksheet.cell(
            row=row_index,
            column=_OVERDUE_AMOUNT_COLUMN_INDEX,
        ).number_format = "₡ #,##0.00"
